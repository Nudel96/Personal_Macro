"""Independent openpyxl/Decimal audit of JST R6 and native public Atlas snapshots.

Use --native for the explicit Rust HTTP test output, or --cache for the one
allowlisted isolated Tauri profile. No production journal or database is read.
"""
import argparse
import datetime
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import sqlite3
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[4]
PROFILE = "com.personal-macro.atlas-jst-20260909"


def read(path):
    return json.loads(path.read_text(encoding="utf-8-sig"), parse_float=Decimal)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT / "apps/desktop/.tmp/atlas-validation/macrohistory/raw/JSTdatasetR6.xlsx")
    parser.add_argument("--native", type=Path)
    parser.add_argument("--cache", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    cfg = read(ROOT / "apps/desktop/src/features/world-atlas/data/macrohistory-catalog.json")
    sha = hashlib.sha256(args.source.read_bytes()).hexdigest()
    assert sha == cfg["sha256"]
    workbook = load_workbook(args.source, read_only=True, data_only=True)
    assert workbook.sheetnames == ["Sheet1"]
    rows = workbook.active.iter_rows(values_only=True)
    headers = list(next(rows))
    assert headers == cfg["headers"] and len(headers) == 59
    source = {}
    identities = set()
    for values in rows:
        row = dict(zip(headers, values))
        key = (row["iso"], row["year"])
        assert key not in source and 1870 <= row["year"] <= 2020
        source[key] = row
        identities.add((row["iso"], row["country"], row["ifs"]))
    workbook.close()
    assert len(source) == 2718 and len(identities) == 18
    assert identities == {(a["code"], a["providerLabel"], a["ifs"]) for a in cfg["areas"]}
    geographies = {g["iso3"]:g["id"] for g in read(ROOT / "apps/desktop/src/features/world-atlas/data/catalog.json")["geographies"] if g["iso3"]}
    assert all(a["geographyId"] == geographies[a["code"]] for a in cfg["areas"])
    native = None
    if args.native:
        native = {r["profile"]["providerCode"]:r for r in read(args.native)}
    if args.cache:
        assert not args.native
        path = args.cache.resolve(strict=True)
        assert PROFILE in path.parts and path.name == "cache.sqlite" and path.parent.name == "atlas"
        with sqlite3.connect(path.as_uri()+"?mode=ro",uri=True) as db:
            provenance = json.loads(db.execute("SELECT provenance_json FROM atlas_macrohistory_dataset WHERE id=?",(cfg["datasetId"],)).fetchone()[0],parse_float=Decimal)
            native = {}
            for (serialized,) in db.execute("SELECT profile_json FROM atlas_macrohistory_areas WHERE dataset_id=?",(cfg["datasetId"],)):
                p=json.loads(serialized,parse_float=Decimal)
                native[p["providerCode"]]={"profile":p,"provenance":provenance,"status":"available"}
    if native is not None:
        assert set(native)=={x[0] for x in identities}
    # These formula definitions are independent of the application's recipe dispatcher.
    direct = {"output":"rgdpbarro","consumption":"rconsbarro","unemployment":"unemp","bankCapital":"lev","loanDeposits":"ltd","noncore":"noncore","shortRate":"stir","longRate":"ltrate"}
    percent = {"publicDebt":"debtgdp","dividendYield":"eq_dp","rentYield":"housing_rent_yd"}
    ratios = {"bankCredit":("tloans","gdp"),"mortgages":("tmort","gdp"),"businessDebt":("bdebt","gdp"),"housePrices":("hpnom","cpi"),"wages":("wage","cpi")}
    returns = {"equityReturn":"eq_tr","housingReturn":"housing_tr","bondReturn":"bond_tr","billReturn":"bill_rate"}
    flags = {"equityReturn":["eq_tr_interp"],"housingReturn":["rent_ipolated","housing_capgain_ipolated"],"housePrices":["housing_capgain_ipolated"],"dividendYield":["eq_dp_interp"],"rentYield":["rent_ipolated"]}
    assert set(direct)|set(percent)|set(ratios)|set(returns)=={m["id"] for m in cfg["metrics"]}
    counts={"rawNumbers":0,"rawMissing":0,"rawZeros":0,"derivedNumbers":0,"derivedMissing":0,"nominalReturns":0,"sourceInterpolationFlags":0,"crisisStarts":0,"comparedNumericCells":0,"exactDecimalMatches":0}
    maximum_relative=Decimal(0)
    coverage={}

    def decimal(n):
        return None if n is None else Decimal(str(n))

    def check(actual,wanted):
        nonlocal maximum_relative
        assert (actual is None)==(wanted is None),(actual,wanted)
        if wanted is None:return
        actual=decimal(actual); wanted=decimal(wanted)
        error=abs(actual-wanted)
        assert error<=max(Decimal("1e-11"),abs(wanted)*Decimal("1e-13")),(actual,wanted,error)
        maximum_relative=max(maximum_relative,error/max(Decimal(1),abs(wanted)))
        counts["comparedNumericCells"]+=1
        counts["exactDecimalMatches"]+=int(error==0)

    for code,label,ifs in sorted(identities):
        coverage[code]={"geographyId":geographies[code],"providerLabel":label,"ifs":ifs,"metrics":{}}
        cache={}
        if native is not None:
            row=native[code]
            assert row["status"]=="available" and row["profile"]["geographyId"]==geographies[code]
            assert row["provenance"]["sha256"]==sha and row["provenance"]["recipe"]==cfg["recipe"]
            cache={p["year"]:p for p in row["profile"]["points"]}
            assert len(cache)==151 and set(cache)==set(range(1870,2021))
        for year in range(1870,2021):
            row=source[code,year]
            raw={k:decimal(v) for k,v in row.items() if k not in {"year","country","iso","ifs","peg_type","peg_base"}}
            assert len(raw)==53
            counts["crisisStarts"]+=int(raw["crisisJST"]==1)
            for key,wanted in raw.items():
                counts["rawMissing" if wanted is None else "rawNumbers"]+=1
                counts["rawZeros"]+=int(wanted==0)
                if native is not None: check(cache[year]["raw"][key],wanted)
            expected={id:raw[field] for id,field in direct.items()}
            expected.update({id:None if raw[field] is None else raw[field]*100 for id,field in percent.items()})
            expected.update({id:None if raw[a] is None or raw[b] is None or raw[b]<=0 else raw[a]/raw[b]*100 for id,(a,b) in ratios.items()})
            nominal={}
            prior=decimal(source.get((code,year-1),{}).get("cpi"))
            for id,field in returns.items():
                nominal[id]=None if raw[field] is None else raw[field]*100
                expected[id]=None if raw[field] is None or raw["cpi"] is None or raw["cpi"]<=0 or prior is None or prior<=0 else ((1+raw[field])*prior/raw["cpi"]-1)*100
                counts["nominalReturns"]+=int(nominal[id] is not None)
            for id,wanted in expected.items():
                counts["derivedMissing" if wanted is None else "derivedNumbers"]+=1
                interpolated=any(raw[key]==1 for key in flags.get(id,[]))
                counts["sourceInterpolationFlags"]+=int(interpolated)
                if wanted is not None:
                    entry=coverage[code]["metrics"].setdefault(id,{"n":0,"first":year,"last":year})
                    entry["n"]+=1; entry["last"]=year
                if native is not None:
                    got=cache[year]["values"][id]
                    check(got["value"],wanted); check(got["nominal"],nominal.get(id))
                    assert got["interpolated"]==interpolated
    result={"checkedAt":datetime.datetime.now(datetime.timezone.utc).isoformat(),"status":"native_check_passed" if native is not None else "source_review_passed","sourceUrl":cfg["url"],"resolvedUrl":cfg["resolvedUrl"],"sha256":sha,"sourceRows":2718,"areas":18,"perspectives":20,"counts":counts,"maximumRelativeRepresentationDifference":str(maximum_relative),"coverage":coverage,"limits":["JST historical reconstructions end in 2020 and cover 18 countries, not the world.","Source definitions and territorial boundaries change; no complete observation-level break list is available.","Real returns require a previous-year CPI value; source nominal returns and interpolation flags are retained.","The audit verifies transcription and stated formulas, not every historical source study or a natural cycle."]}
    args.output.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({k:v for k,v in result.items() if k not in {"coverage","limits"}},ensure_ascii=True))


if __name__=="__main__":main()
