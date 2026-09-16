"""Audit published EPO chart caches against their embedded original workbooks.

No chart digitisation, patent searches, contemporary backcasts or country sums.
Only the years covered by the final reports are used; bubble sizes are counts,
and their artificial x coordinates are resolved through the published year labels.
"""
import hashlib
import json
import posixpath
import re
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl

ROOT = Path(__file__).resolve().parents[4]
HERE = Path(__file__).resolve().parent
DATA = ROOT / "apps/desktop/src/features/world-atlas/data"
RAW = ROOT / "apps/desktop/.tmp/atlas-remaining-40"
NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
VERSION = "2026-09-11.17"


def read(p):
    return json.loads(p.read_text(encoding="utf-8"))


def write(p, obj):
    p.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def cache(series, tag):
    return {int(p.attrib["idx"]): p.find("c:v", NS).text
            for p in series.findall(f"c:{tag}//c:pt", NS)}


def formula(series, tag):
    f = series.find(f"c:{tag}//c:f", NS).text
    m = re.fullmatch(r"([^!]+)!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)", f)
    assert m and m[2] == m[4], f
    return m[1], openpyxl.utils.column_index_from_string(m[2]) - 1, int(m[3]) - 1, int(m[5]) - 1


def integer(n):
    assert type(n) in (int, float) and n == int(n) and 0 <= n < 100000, n
    return str(int(n))


CATALOG = read(DATA / "catalog.json")
GEO = {g["iso3"]: g for g in CATALOG["geographies"] if g.get("iso3")}
# These are source identities, not inferred from filing authority prefixes.
ORIGINS = {
    "Germany": "DEU", "France": "FRA", "United Kingdom": "GBR", "Italy": "ITA",
    "Sweden": "SWE", "Spain": "ESP", "Netherlands": "NLD", "Switzerland": "CHE",
    "Austria": "AUT", "Belgium": "BEL", "Norway": "NOR", "Luxembourg": "LUX",
    "Finland": "FIN", "Greece": "GRC", "Turkey": "TUR", "Romania": "ROU",
    "Ireland": "IRL", "Portugal": "PRT", "Cyprus": "CYP", "Denmark": "DNK",
    "Poland": "POL", "Czech Republic": "CZE", "Lithuania": "LTU", "Hungary": "HUN",
}
QUANTUM_NOTE = (
    "Archivstudie 2019 über Quantensensorik und Quantenmesstechnik der zweiten Generation. "
    "Keine vollständige Erfassung von Quantencomputern oder Quantenkommunikation. "
    "Jahre bezeichnen Veröffentlichungen, nicht Umsätze oder einen aktuellen Branchenzustand. "
    "Nur 2000–2017 aus dem endgültigen Bericht; die zusätzliche Präsentationsspalte 2018 bleibt ausgeschlossen. "
    "Patentfamilien und Prioritätsanmeldungen sind unterschiedliche Zählungen; Länder werden nicht zur Welt summiert. "
    "Suchstrategie und Veröffentlichungsverzug begrenzen die Abdeckung. Quelle: Europäisches Patentamt, 2019. "
    "Deutsche Erläuterung und Diagramm durch Personal Macro."
)
SPACE_NOTE = (
    "Archivstudie 2021: Cosmonautics, überwiegend CPC B64G und 41 Suchfelder des ESA-Technologiebaums. "
    "Die Recherche endete im November 2019; die Atlasgeschichte endet 2017 wegen des Veröffentlichungsverzugs. "
    "Startsysteme sind außerhalb dieser Patentabgrenzung. Familien zählen verwandte Anmeldungen einer Erfindung; "
    "Anmeldungen zählen Schutzanträge und sind keine Erfindungs-, Produktions- oder Umsatzstatistik. "
    "EPO38+ ist die feste Schutzgebietsabgrenzung der Studie von 2019, kein heutiges EU-Aggregat. "
    "Fehlende Diagrammpunkte werden nicht zu Null. Quelle: EPO/ESPI/ESA, Cosmonautics 2021. "
    "Deutsche Erläuterung und Diagramm durch Personal Macro."
)


def build(kind):
    quantum = kind == "quantum"
    sid = "epo-quantum-sensing" if quantum else "epo-cosmonautics"
    file = "epo-quantum-sensing.pptx" if quantum else "epo-cosmonautics.pptx"
    url = "https://link.epo.org/web/" + ("insights-quantum_metrology_and_sensing_en.pptx" if quantum else "data_mapping_of_cosmonautics_graphs_and_datasets_en.pptx")
    document = "https://link.epo.org/web/" + ("patent_insight_report-quantum_metrology_and_sensing_en.pdf" if quantum else "patent_insight_report-cosmonautics_en.pdf")
    first = 2000 if quantum else 1990
    note = QUANTUM_NOTE if quantum else SPACE_NOTE
    raw = (RAW / file).read_bytes()
    z = ZipFile(BytesIO(raw))
    books, columns, profiles, metrics, areas = [], [], [], [], {}

    def metric(code, label, unit, explanation, scope):
        item = {"id": f"{sid}:{code}", "sourceId": sid,
                "topicId": "innovation:quantum" if quantum else "innovation:space_technology",
                "providerCode": code, "label": label, "unit": unit, "frequency": "annual",
                "kind": "source_statistic", "comparison": "same_definition", "connectAdjacent": True,
                "explanation": explanation, "scopeNote": scope + " " + note}
        metrics.append(item)
        return item

    def book(chart_no):
        chart_path = f"ppt/charts/chart{chart_no}.xml"
        rels = ET.fromstring(z.read(f"ppt/charts/_rels/chart{chart_no}.xml.rels"))
        member = posixpath.normpath(posixpath.join("ppt/charts", next(r.attrib["Target"] for r in rels if r.attrib["Type"].endswith("/package"))))
        content = z.read(member)
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
        spec = {"member": member, "sha256": hashlib.sha256(content).hexdigest(),
                "sheetNames": workbook.sheetnames, "headers": [], "axisColumn": None, "axisFirstRow": None,
                "chart": chart_path, "chartSha256": hashlib.sha256(z.read(chart_path)).hexdigest()}
        books.append(spec)
        return len(books) - 1, workbook, ET.fromstring(z.read(chart_path)).findall(".//c:ser", NS), spec

    def add_profile(metric, area_code, area_label, iso3, book_index, workbook, series, value_tag,
                    title, column_note, first_year=first, axis=None):
        sheet, value_col, row_first, row_last = formula(series, value_tag)
        year_sheet, year_col, yr_first, yr_last = formula(series, "cat" if axis is None else "xVal")
        assert sheet == year_sheet and (row_first, row_last) == (yr_first, yr_last)
        ws = workbook[sheet]
        by_index, year_cache = cache(series, value_tag), cache(series, "cat" if axis is None else "xVal")
        values = {}
        for index in range(row_last - row_first + 1):
            cell_year = ws.cell(row_first + index + 1, year_col + 1).value
            actual_year = int(cell_year) if axis is None else int(axis[int(cell_year) - 1])
            assert int(year_cache[index]) == int(cell_year)
            raw_value = ws.cell(row_first + index + 1, value_col + 1).value
            assert integer(raw_value) == by_index[index]
            assert actual_year not in values
            if first_year <= actual_year <= 2017:
                values[actual_year] = by_index[index]
        assert values
        # Validate the full original column label, including those shortened in chart legends.
        header_column = value_col if axis is None else value_col - 1
        header = ws.cell(1, header_column + 1).value
        books[book_index]["headers"].append({"sheet": sheet, "row": 0, "column": header_column, "value": str(header)})
        if axis is not None:
            assert ws.cell(1, value_col + 1).value == str(header) + " - Size"
            books[book_index]["headers"].append({"sheet": sheet, "row": 0, "column": value_col, "value": str(header) + " - Size"})
        geo = GEO[iso3]["id"]
        area = areas.setdefault(area_code, {"code": area_code, "label": area_label, "geographyId": geo, "seriesTitles": {}})
        area["seriesTitles"][metric["providerCode"]] = title
        columns.append({"metricCode": metric["providerCode"], "areaCode": area_code, "areaLabel": area_label,
                        "geographyId": geo, "book": book_index, "sheet": sheet, "yearColumn": year_col,
                        "valueColumn": value_col, "firstRow": row_first, "lastRow": row_last,
                        "firstYear": first_year, "lastYear": 2017, "title": title, "note": column_note})
        points = [{"period": str(y), "value": values.get(y),
                   "status": "Archivstudie 2019" if quantum else "Archivstudie 2021",
                   "breakBefore": False, "notes": [column_note] + ([] if y in values else ["Kein veröffentlichter Diagrammpunkt; nicht als Null gezählt."]),
                   "lowerBound": None, "upperBound": None} for y in range(first_year, 2018)]
        profiles.append({"metricId": metric["id"], "geographyId": geo, "providerArea": area_code,
                         "providerLabel": area_label, "providerTitle": title, "unit": metric["unit"], "points": points})

    if quantum:
        m = metric("families", "Quantensensorik · weltweite Patentfamilien", "Patentfamilien (Anzahl)",
                   "Wie sich veröffentlichte Patentfamilien zur Quantensensorik in der historischen EPA-Recherche entwickeln.",
                   "Weltweite Suchergebnisse, geordnet nach dem ersten Veröffentlichungsjahr.")
        idx, w, ss, spec = book(3)
        add_profile(m, "WLD", "Worldwide", "WLD", idx, w, ss[0], "val",
                    "Number of patent families by first publication year", "Weltweite Patentfamilien nach erstem Veröffentlichungsjahr; historische Quantensensorik-Recherche.")
        m = metric("priority", "Quantensensorik · Prioritätsanmeldungen", "Prioritätsanmeldungen (Anzahl)",
                   "Wie viele Erstanmeldungen die Studie einer Patentzuständigkeit und einem Veröffentlichungsjahr zuordnet.",
                   "CN, US und JP sind Zuständigkeiten der ersten Anmeldung. Daraus wird kein sicherer Wohnsitz von Erfindern oder Firmen abgeleitet.")
        idx, w, ss, spec = book(4)
        for i, (code, label, iso3) in enumerate([("CN", "China", "CHN"), ("US", "United States", "USA"), ("JP", "Japan", "JPN")]):
            assert cache(ss[i], "tx")[0] == code
            add_profile(m, code, label, iso3, idx, w, ss[i], "val",
                        "Number of priority filings by jurisdiction and publication year | " + code,
                        "Erste Patentzuständigkeit " + code + "; Veröffentlichungsjahr. Kein belegter Wohnsitz von Erfindern oder Unternehmen.")
    else:
        for chart_no, suffix, territory in [(4, "world", "weltweit"), (5, "epo38", "Schutz in EPO38+")]:
            idx, w, ss, spec = book(chart_no)
            for i, (code, noun, unit) in enumerate([("families", "Patentfamilien", "Patentfamilien (Anzahl)"), ("applications", "Anmeldungen", "Patentanmeldungen (Anzahl)")]):
                m = metric(code + "_" + suffix, f"Raumfahrt · {noun} · {territory}", unit,
                           f"Wie sich die veröffentlichten {noun} der Cosmonautics-Recherche im Bereich {territory} entwickeln.",
                           "Weltweiter Quellenbestand." if suffix == "world" else "Weltweite Anmelder mit Schutz in der festen EPO38+-Abgrenzung von 2019; kein Herkunftsbild europäischer Erfinder.")
                title = f"Cosmonautics – {'worldwide patent filing developments' if suffix == 'world' else 'developments in patents filed in the EPO 38+, not limited to appl. from Europe'} | " + cache(ss[i], "tx")[0]
                add_profile(m, "WLD", "Worldwide", "WLD", idx, w, ss[i], "val", title,
                            "Anmeldejahr; " + ("weltweite Patentaktivität." if suffix == "world" else "Schutz in EPO38+ laut fester Gebietsabgrenzung 2019, unabhängig vom Herkunftsland."))
        idx, w, ss, spec = book(11)
        axis = [w["ChartData"].cell(row, 67).value for row in range(511, 541)]
        assert axis == [str(y) for y in range(1990, 2020)]
        spec.update(axisColumn=66, axisFirstRow=510)
        m = metric("applicant_origin", "Raumfahrt · Anmeldungen nach Herkunft", "Patentanmeldungen (Anzahl)",
                   "Wie sich die in der Studie ausgewiesenen Anmeldungen nach dem Herkunftsland des Anmelders entwickeln.",
                   "Figure 12 umfasst Herkunftsländer innerhalb der damaligen EPO38+-Region. Das ist keine vollständige Welttabelle und keine Zuordnung nach Patentamt.")
        for series in ss:
            if "xVal" not in [c.tag.rsplit('}', 1)[-1] for c in series]:
                continue
            sheet, col, row_first, row_last = formula(series, "bubbleSize")
            label = w[sheet].cell(1, col).value
            if label not in ORIGINS:
                continue
            iso = ORIGINS[label]
            add_profile(m, iso, label, iso, idx, w, series, "bubbleSize",
                        "Cosmonautics – development of filings from EPO38+ by country of applicant | " + label,
                        "Herkunftsland des Anmelders laut Figure 12; Anmeldejahr. Nur veröffentlichte Diagrammpunkte werden gezählt.", axis=axis)

    profiles.sort(key=lambda p: (p["metricId"], p["geographyId"]))
    source = {"id": sid, "label": "EPA · Quantensensorik · Archiv 2019" if quantum else "EPA / ESPI / ESA · Raumfahrtpatente · Archiv 2021",
              "adapter": "epo_embedded", "url": url, "documentationUrl": document,
              "licenseUrl": "https://www.epo.org/en/terms-of-use", "publishedAt": "2019-09" if quantum else "2021-07",
              "reviewedAt": "2026-09-11", "recipe": sid + "-original-chart-workbooks-through2017-v1",
              "observationKind": "source_statistics", "expectedSha256": hashlib.sha256(raw).hexdigest(),
              "expectedRows": sum(len(p["points"]) for p in profiles),
              "expectedNumeric": sum(pt["value"] is not None for p in profiles for pt in p["points"]),
              "firstPeriod": str(first), "lastPeriod": "2017", "areas": sorted(areas.values(), key=lambda a: a["code"])}
    contract = {"sourceId": sid, "file": file, "zipEntries": len(z.infolist()), "firstYear": first, "lastYear": 2017,
                "books": books, "columns": columns}
    audit = {"source": source, "profileCount": len(profiles), "numericValues": source["expectedNumeric"],
             "missingPoints": source["expectedRows"] - source["expectedNumeric"], "contract": contract,
             "independentCheck": "Every selected numeric chart-cache value equals the embedded workbook, including zero, source year coordinates and labels.",
             "rawFileBytes": len(raw), "uncompressedBytes": sum(e.file_size for e in z.infolist()),
             "excluded": ["2018/2019 pending publication years are excluded; no later years reconstructed.",
                          "No office-to-inventor-origin conversion. EPO/PCT procedural areas are not assigned to countries or a continent.",
                          "Former Serbia and Montenegro is not assigned to successor states.",
                          "Space launch, commercial revenues, generic patents and private applicant rankings are not imported."],
             "newerResearch": [
                 {"source": "EPO/OECD Mapping the global quantum ecosystem 2025", "finding": "Newer report reviewed. No accompanying numeric history found. The 2023 computing/simulation Excel downloads contain search strategies, not observations; not spliced into the 2019 archive."},
                 {"source": "EPO Space propulsion 2024/2025", "finding": "Actual 72 MB master file examined; it has application, applicant and technology duplicates requiring a separate family-count methodology. It is not substituted for the published Cosmonautics series."},
                 {"source": "OECD Space Economy 2026", "finding": "New report and supplementary materials examined. No new data series are claimed from publication text alone."}]}
    write(RAW / f"{sid}-expected-profiles.json", profiles)
    return source, metrics, contract, audit


results = [build("quantum"), build("space")]
ids = {r[0]["id"] for r in results}
public = read(DATA / "public-series-catalog.json")
public["sources"] = [s for s in public["sources"] if s["id"] not in ids] + [r[0] for r in results]
public["metrics"] = [m for m in public["metrics"] if m["sourceId"] not in ids] + [m for r in results for m in r[1]]
public["version"] = VERSION
write(DATA / "public-series-catalog.json", public)
write(DATA / "public-epo-contract.json", [r[2] for r in results])
write(HERE / "public-epo-audit.json", [r[3] for r in results])
CATALOG["version"] = VERSION
write(DATA / "catalog.json", CATALOG)
ledger = read(HERE / "remaining-40-ledger.json")
for result, topic_id in zip(results, ["innovation:quantum", "innovation:space_technology"]):
    t = next(t for t in ledger["topics"] if t["id"] == topic_id)
    t["research"] = [{"sourceId": result[0]["id"], "reviewedAt": "2026-09-11", "evidence": "public-epo-audit.json"}]
    if t["status"] != "implemented_native_verified":
        t["status"] = "source_validated"
write(HERE / "remaining-40-ledger.json", ledger)
print(json.dumps([{ "source": r[0]["id"], "areas": len(r[0]["areas"]), "profiles": r[3]["profileCount"], "values": r[3]["numericValues"], "missing": r[3]["missingPoints"]} for r in results]))
