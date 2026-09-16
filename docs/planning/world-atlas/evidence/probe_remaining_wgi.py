"""Inspect the publisher-linked WGI 2025 revision workbook without changing user data."""
import hashlib
import json
import urllib.request
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[4]
DEST = ROOT / "apps/desktop/.tmp/atlas-remaining-40"
URL = "https://datacatalogfiles.worldbank.org/ddh-published/0038026/DR0095947/wgidataset_with_sourcedata-2025.xlsx"
path = DEST / "WGI2025.xlsx"
if not path.exists():
    req = urllib.request.Request(URL, headers={"User-Agent": "PersonalMacro-Atlas source research"})
    with urllib.request.urlopen(req, timeout=90) as response:
        data = response.read(32 * 1024 * 1024 + 1)
        assert len(data) <= 32 * 1024 * 1024
    path.write_bytes(data)
data = path.read_bytes()
book = openpyxl.load_workbook(path, read_only=True, data_only=True)
out = {"url": URL, "sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data), "sheets": {}}
for sheet in book:
    first = []
    for i, row in enumerate(sheet.values):
        if i < 12:
            first.append(list(row))
        else:
            break
    out["sheets"][sheet.title] = {"rows": sheet.max_row, "cols": sheet.max_column, "firstRows": first}
(DEST / "WGI2025-inspection.json").write_text(json.dumps(out, ensure_ascii=False, indent=2, default=str), encoding="utf8")
print(json.dumps({"bytes": out["bytes"], "sha256": out["sha256"], "sheets": {k: {"rows": v["rows"], "cols": v["cols"]} for k, v in out["sheets"].items()}}, indent=2))
