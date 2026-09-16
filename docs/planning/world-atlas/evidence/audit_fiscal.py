"""Independent openpyxl/Decimal check of IMF Dec 2025, identities, original values and scope flags.

No user journal is read. --cache only accepts the dedicated isolated test profile.
"""
import argparse
import csv
from collections import Counter
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import sqlite3
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[4]
PROFILE='com.personal-macro.atlas-fiscal-20260909'

def read(path):
    return json.loads(path.read_text(encoding='utf-8-sig'),parse_float=Decimal)

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--source',type=Path,default=ROOT/'.tmp/atlas-validation/fiscal/snapshot.xlsx')
    p.add_argument('--native',type=Path)
    p.add_argument('--cache',type=Path)
    p.add_argument('--output',type=Path,required=True)
    args=p.parse_args()
    cfg=read(ROOT/'apps/desktop/src/features/world-atlas/data/fiscal-catalog.json')
    raw=args.source.read_bytes();sha=hashlib.sha256(raw).hexdigest()
    assert sha==cfg['sha256'] and hashlib.md5(raw).hexdigest()==cfg['md5'] and len(raw)==cfg['fileSize']
    w=load_workbook(args.source,read_only=True,data_only=True)
    assert w.sheetnames==['data']
    it=w.active.iter_rows(values_only=True);headers=list(next(it));assert headers==cfg['headers']
    source={}; identities=set()
    for vals in it:
        r=dict(zip(headers,vals));key=(r['isocode'],r['year']);assert key not in source
        source[key]=r;identities.add((r['country'],r['isocode'],r['ifscode']))
    w.close()
    assert len(source)==33846 and len(identities)==153
    assert identities=={(label,a['code'],a['ifs']) for a in cfg['areas'] for label in a['providerLabels']}
    geographies={a['iso3']:a['id'] for a in read(ROOT/'apps/desktop/src/features/world-atlas/data/catalog.json')['geographies'] if a['iso3'] and a['kind']!='aggregate'}
    native=None
    if args.native:native={r['profile']['providerCode']:r for r in read(args.native)}
    if args.cache:
        assert not args.native
        path=args.cache.resolve(strict=True)
        assert PROFILE in path.parts and path.name=='cache.sqlite' and path.parent.name=='atlas'
        with sqlite3.connect(path.as_uri()+'?mode=ro',uri=True) as db:
            prov=json.loads(db.execute('SELECT provenance_json FROM atlas_fiscal_dataset WHERE id=?',(cfg['datasetId'],)).fetchone()[0],parse_float=Decimal)
            native={}
            for (s,) in db.execute('SELECT profile_json FROM atlas_fiscal_areas WHERE dataset_id=?',(cfg['datasetId'],)):
                profile=json.loads(s,parse_float=Decimal);native[profile['providerCode']]={'profile':profile,'provenance':prov,'status':'available'}
    if native is not None:assert set(native)=={a['code'] for a in cfg['areas']}
    counts=Counter();coverage={};maximum=Decimal(0)
    for area in cfg['areas']:
        code=area['code'];assert geographies[code]==area['geographyId']
        observations={year:r for (c,year),r in source.items() if c==code}
        assert len(observations)==area['rowCount'] and set(observations)==set(range(area['firstRowYear'],area['lastRowYear']+1))
        destination={}
        if native is not None:
            n=native[code];assert n['status']=='available' and n['profile']['geographyId']==area['geographyId']
            assert n['profile']['providerLabels']==area['providerLabels']
            assert n['provenance']['sha256']==sha and n['provenance']['recipe']==cfg['recipe']
            assert n['provenance']['url']==cfg['url'] and n['provenance']['originalUrl']==cfg['originalUrl']
            destination={p['year']:p for p in n['profile']['points']};assert set(destination)==set(observations)
        coverage[code]={'geographyId':area['geographyId'],'providerLabels':area['providerLabels'],'ifs':area['ifs'],'metrics':{}}
        for year,r in observations.items():
            for key,dkey in [('GG_budg','budgetScope'),('GG_debt','debtScope')]:
                assert r[key] in (0,1);counts['scopeFlags']+=1
                if native is not None:assert destination[year][dkey]==r[key]
            for key in headers[6:]:
                value=r[key];counts['missing' if value is None else 'numeric']+=1
                counts['zeros']+=int(value==0);counts['negative']+=int(value is not None and value<0)
                if native is not None:
                    found=destination[year]['values'][key];assert (found is None)==(value is None)
                    if value is not None:
                        wanted=Decimal(str(value));err=abs(found-wanted);maximum=max(maximum,err/max(1,abs(wanted)))
                        assert err<=max(Decimal('1e-12'),abs(wanted)*Decimal('1e-13')),(code,year,key,found,wanted)
                        counts['comparedNumericCells']+=1;counts['exactDecimalMatches']+=int(err==0)
        for m in cfg['metrics']:
            valid=[r for r in observations.values() if r[m['field']] is not None];years=sorted(r['year'] for r in valid)
            scope=m['scopeField'];scope_counts=Counter(r[scope] for r in valid) if scope else Counter()
            coverage[code]['metrics'][m['id']]={'n':len(years),'first':min(years,default=None),'last':max(years,default=None),'governmentScopes':dict(scope_counts)}
    # Public OWID Grapher CSV is a second, rounded publication of this release.
    csv_path=ROOT/'.tmp/atlas-validation/fiscal/spending.csv'
    if csv_path.exists():
        for row in csv.DictReader(csv_path.open(encoding='utf-8-sig',newline='')):
            actual=Decimal(row['Government expenditure (% of GDP)']);original=Decimal(str(source[row['Code'],int(row['Year'])]['exp']))
            assert abs(actual-original)<=max(Decimal('0.0001'),abs(original)*Decimal('1e-7'))
            counts['owidRoundedExpenditureChecks']+=1
        assert counts['owidRoundedExpenditureChecks']==9265
    report={'checkedAt':'2026-09-09','source':'IMF Public Finances in Modern History, December 2025; unmodified workbook via public OWID snapshot','url':cfg['url'],'originalUrl':cfg['originalUrl'],'sha256':sha,'bytes':len(raw),'profiles':151,'originalLabelIdentities':153,'rows':33846,'checks':dict(counts),'maximumRelativeNumericDifference':str(maximum),'nativeVerified':native is not None,'productionDataAccessed':False,'coverage':coverage,'limits':['Reconstruction, calendar/fiscal years and historical territorial definitions are not fully marked.','Scope flags denote central/general government only where the corresponding value exists.','IMF appendix February 2025 and DataMapper advertised counts differ from this frozen December 2025 workbook.','No world or continent aggregate. No fair values, debt sustainability score or inferred sine cycle.']}
    args.output.parent.mkdir(parents=True,exist_ok=True);args.output.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({k:v for k,v in report.items() if k!='coverage'},ensure_ascii=False))

if __name__=='__main__':main()
