"""Independently audit original Sasol operating tables, never equity prices or forecasts."""
import hashlib
import json
import re
import zipfile
import xml.etree.ElementTree as ET
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[4]
HERE=Path(__file__).resolve().parent
DATA=ROOT/'apps/desktop/src/features/world-atlas/data'
RAW=ROOT/'apps/desktop/.tmp/atlas-remaining-40'
ID='sasol-synthetic-fuels'
VERSION='2026-09-11.22'
DOC='https://www.sasol.com/investor-centre/financial-results'
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def write(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def norm(t):return ' '.join(t.split())

definitions=[
 ('secunda_total','Secunda · raffinierte Syntheseprodukte gesamt','ZA','Millionen Barrel je Geschäftsjahr',
  'Veröffentlichte Gesamtmenge raffinierter Syntheseprodukte der Secunda Operations. Die ältere, ausdrücklich auf White Product begrenzte Reihe bleibt eine eigene Perspektive.'),
 ('secunda_white','Secunda · White Product · Archiv 2014–2019','ZA','Millionen Barrel je Geschäftsjahr',
  'Engere historische Reihe Synfuels refined product (white product). Keine Verlängerung mit der ab 2020 ausgewiesenen Gesamtmenge raffinierter Produkte.'),
 ('oryx_production','ORYX GTL · von Sasol berichtete Produktion','QA','Millionen Barrel je Geschäftsjahr',
  'Veröffentlichte ORYX-GTL-Produktionsmenge aus Sasols Beteiligungsberichterstattung. Sasol hält 49 Prozent an der Anlage. Die berichteten Mengen werden nicht auf die gesamte Anlage hochgerechnet.'),
 ('oryx_utilisation','ORYX GTL · Auslastung der Nennkapazität','QA','Auslastung der Nennkapazität (%)',
  'Originale jährliche Kapazitätsauslastung der ORYX-GTL-Anlage; 2014–2023. Prozentangaben werden nicht mit einer Nennkapazität multipliziert, um fehlende Produktionsmengen zu schätzen.'),
]
files=[];all_values={};raw_cells=0
def add_value(code,year,value,filename,location):
 global raw_cells
 assert code in {d[0] for d in definitions} and 2014<=year<=2026
 n=Decimal(value);assert n.is_finite() and n>=0 and n<=Decimal(150 if code=='oryx_utilisation' else 100)
 key=(code,year)
 if key in all_values:assert Decimal(all_values[key]['value'])==n,(key,all_values[key],value)
 all_values[key]={'value':value,'file':filename,'location':location};raw_cells+=1

pdfs=[(2016,'https://www.sasol.com/sites/default/files/2023-11/Analyst%20Book%20FY16_0.pdf',20,[2016,2015,2014]),
      (2019,'https://sasol.com/sites/default/files/2023-11/Additional%20Analyst%20information%20for%2030%20June%202019.pdf',41,[2019,2018,2017])]
for release,url,page,years in pdfs:
 name=f'sasol-analysts-{release}.pdf';b=(RAW/name).read_bytes();pdf=PdfReader(RAW/name);t=norm(pdf.pages[page-1].extract_text())
 assert 'Energy*' in t and 'Synfuels refined product (white product)' in t and 'ORYX GTL' in t
 tail=t[t.index('ORYX GTL'):]
 numeric=r'([0-9]+,[0-9]+)'
 change=r'(?:[–−-]|\(?[0-9]+\)?)\s+' if release==2019 else ''
 specs=[('secunda_white',r'Synfuels refined product \(white product\)\s*[³⁴34]?\s+mm bbl\s+'+change+numeric+r'\s+'+numeric+r'\s+'+numeric,t),
        ('oryx_production',r'Production\s+mm bbl\s+'+change+numeric+r'\s+'+numeric+r'\s+'+numeric,tail),
        ('oryx_utilisation',r'Utilisation rate of nameplate capacity(?:\s+[–-]\s+ORYX GTL\s*[⁴4]?)?\s+%\s+([0-9]+)\s+([0-9]+)\s+([0-9]+)',tail)]
 rows=[]
 for code,pattern,section in specs:
  m=re.findall(pattern,section);assert len(m)==1,(name,code,m);values=[v.replace(',','.') for v in m[0]]
  rows.append({'code':code,'pattern':pattern,'afterOryx':code.startswith('oryx'),'values':values})
  for year,value in zip(years,values):add_value(code,year,value,name,f'PDF-Seite {page}')
 files.append({'file':name,'url':url,'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest(),'format':'pdf','page':page,'pages':len(pdf.pages),'years':years,'rows':rows})

xlsx_specs=[
 (2022,'https://www.sasol.com/sites/default/files/2023-11/Production%20and%20sales%20metrics%20for%20the%20year%20ended%2030%20June%202022_Excel%20data_0.xlsx',[6,7,8],[2022,2021,2020],4,[(13,'secunda_total',2,'Synfuels total refined product'),(20,'oryx_production',3,'Production'),(21,'oryx_utilisation',3,'Utilisation rate of nameplate capacity')]),
 (2023,'https://www.sasol.com/sites/default/files/2023-11/Production%20and%20sales%20metrics%20for%20the%20year%20ended%2030%20June%202023%20-%20Excel%20Data.xlsx',[5,6,7],[2023,2022,2021],3,[(11,'secunda_total',1,'Secunda Operations total refined product'),(18,'oryx_production',2,'Production'),(19,'oryx_utilisation',2,'Utilisation rate of nameplate capacity')]),
 (2025,'https://www.sasol.com/sites/default/files/2025-07/Production%20and%20sales%20metrics%20for%20the%20year%20ended%2030%20June%202025_excel.xlsx',[8,9,10],[2025,2024,2023],3,[(9,'secunda_total',1,'Secunda Operations total refined'),(11,'oryx_production',1,'ORYX GTL production')]),
 (2026,'https://www.sasol.com/sites/default/files/2026-07/Business%20Performance%20Metrics%20for%20the%20year%20ended%2030%20June%202026_excel.xlsx',[8,9,10],[2026,2025,2024],3,[(9,'secunda_total',1,'Secunda Operations total refined'),(11,'oryx_production',1,'ORYX GTL production')])]
ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
for release,url,columns,years,unit_col,specs in xlsx_specs:
 name=f'sasol-production-{release}.xlsx';b=(RAW/name).read_bytes();book=load_workbook(RAW/name,read_only=True,data_only=True);sheet=book['Fuels']
 for col,year in zip(columns,years):assert sheet.cell(3,col).value=='Full year' and int(sheet.cell(4,col).value)==year
 with zipfile.ZipFile(RAW/name) as archive:
  xml=ET.fromstring(archive.read('xl/worksheets/sheet3.xml'))
  cells={c.get('r'):c.find('m:v',ns).text for c in xml.findall('.//m:c',ns) if c.find('m:v',ns) is not None and c.get('t') not in ['s','str']}
  entries=len(archive.infolist());uncompressed=sum(e.file_size for e in archive.infolist())
 rows=[]
 for row,code,label_col,label in specs:
  assert norm(str(sheet.cell(row,label_col).value))==label
  unit='%' if code=='oryx_utilisation' else 'mm bbl';assert sheet.cell(row,unit_col).value==unit
  values=[]
  for col,year in zip(columns,years):
   cell=sheet.cell(row,col);value=cells[cell.coordinate];assert float(value)==cell.value
   values.append(value);add_value(code,year,value,name,'Fuels!'+cell.coordinate)
  rows.append({'row':row,'code':code,'labelColumn':label_col,'label':label,'unit':unit,'values':values})
 files.append({'file':name,'url':url,'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest(),'format':'xlsx','sheet':'Fuels','sheetNames':book.sheetnames,'dimensions':[sheet.max_row,sheet.max_column],'dataXml':'xl/worksheets/sheet3.xml','zipEntries':entries,'uncompressedBytes':uncompressed,'columns':columns,'years':years,'unitColumn':unit_col,'rows':rows});book.close()

assert raw_cells==48 and len(all_values)==36
assert [y for c,y in all_values if c=='oryx_utilisation']==list(range(2016,2013,-1))+list(range(2019,2016,-1))+[2022,2021,2020,2023]
scope=('Sasol, originale Betriebskennzahlen, Geschäftsjahre zum 30. Juni. Ein Jahr bezeichnet den Zeitraum Juli des Vorjahres bis Juni des genannten Jahres. '
 'Secunda in Südafrika nutzt Kohle und Erdgas zur Fischer-Tropsch-Synthese; ORYX GTL in Katar wandelt Erdgas in flüssige Produkte um. '
 'Dies sind fossile Synthesewege, keine erneuerbaren E-Fuels. Die Bilder zeigen einzelne Anlagen beziehungsweise die Beteiligungsberichterstattung eines Betreibers, keine nationale oder weltweite Gesamtproduktion. '
 'Sasols Beteiligung an ORYX beträgt 49 Prozent; berichtete Mengen bleiben unverändert, ohne Hochrechnung auf die gesamte Anlage. '
 'White Product 2014–2019 bleibt von Total Refined ab 2020 getrennt. Chemikalien enthaltene Gesamterzeugung, Rohölraffinerie Natref, Verkäufe, Quartale und Prognosen werden nicht verwendet. '
 'Jeder Punkt nennt seinen Originalbericht; überlappende Angaben sind geprüft. Lücken bleiben offen, unterschiedliche Kennzahlen werden nicht addiert. '
 'Öffentliche Betreiberberichte, eigene Quellenprüfung am 11.09.2026; keine offene Gesamtlizenz behauptet.')
areas=[{'code':'ZA','label':'Secunda Operations, South Africa','geographyId':'m49:710','seriesTitles':{}},
       {'code':'QA','label':'ORYX GTL, Ras Laffan, Qatar — Sasol reporting','geographyId':'m49:634','seriesTitles':{}}]
metrics=[];profiles=[]
for code,label,area,unit,note in definitions:
 title={'secunda_total':'Secunda Operations total refined product','secunda_white':'Synfuels refined product (white product)','oryx_production':'ORYX GTL production — as reported by Sasol','oryx_utilisation':'ORYX GTL utilisation rate of nameplate capacity'}[code]
 m={'id':ID+':'+code,'sourceId':ID,'topicId':'fuels:synthetic_fuels','providerCode':code,'label':label,'unit':unit,'frequency':'fiscal_annual_june','kind':'plant_production','comparison':'within_country','connectAdjacent':True,'explanation':note,'scopeNote':scope+' '+note};metrics.append(m)
 a=next(a for a in areas if a['code']==area);a['seriesTitles'][code]=title
 points=[]
 for (c,year),v in sorted(all_values.items()):
  if c!=code:continue
  f=next(f for f in files if f['file']==v['file'])
  points.append({'period':str(year),'value':v['value'],'status':'Originaler Betreiberbericht','breakBefore':False,'notes':[f'Geschäftsjahr {year-1}/{year}, Ende 30. Juni {year}.',note,f"Quelle: {f['file']} · {v['location']}",'Originaldatei: '+f['url']],'lowerBound':None,'upperBound':None})
 profiles.append({'metricId':m['id'],'geographyId':a['geographyId'],'providerArea':a['code'],'providerLabel':a['label'],'providerTitle':title,'unit':unit,'points':points})
profiles.sort(key=lambda p:(p['metricId'],p['geographyId']))
package=b''.join((RAW/f['file']).read_bytes() for f in files);(RAW/'sasol-synthetic-fuels-originals.bin').write_bytes(package)
source={'id':ID,'label':'Sasol · fossile Synthesekraftstoffe · Anlagenberichte','adapter':'sasol_synfuels','url':files[-1]['url'],'documentationUrl':DOC,'licenseUrl':'https://www.sasol.com/legal-notices','publishedAt':'2026-07','reviewedAt':'2026-09-11','recipe':'sasol-six-original-reports-fy2014to2026-v1','observationKind':'operator_report','expectedSha256':hashlib.sha256(package).hexdigest(),'expectedRows':raw_cells,'expectedNumeric':36,'firstPeriod':'2014','lastPeriod':'2026','areas':areas}
contract={'sourceId':ID,'packageFile':'sasol-synthetic-fuels-originals.bin','files':files,'areas':areas,'profiles':profiles,'scope':scope,'supportingMethodology':{'url':'https://www.sec.gov/Archives/edgar/data/314590/000141057825001910/ssl-20250630x20f.htm','sections':'Item 4.D, Fuels plants and facilities; Integrated Report 2025 printed page 59','boundary':'No oil-equivalent reserve production or chemical-inclusive SO production substituted for refined products.'}}
catalog=read(DATA/'public-series-catalog.json');catalog['sources']=[s for s in catalog['sources'] if s['id']!=ID]+[source];catalog['metrics']=[m for m in catalog['metrics'] if m['sourceId']!=ID]+metrics;catalog['version']=VERSION;write(DATA/'public-series-catalog.json',catalog)
main=read(DATA/'catalog.json');main['version']=VERSION;write(DATA/'catalog.json',main)
write(DATA/'public-synfuels-contract.json',contract);write(RAW/'synfuels-expected-profiles.json',profiles)
write(HERE/'public-synfuels-audit.json',{'reviewedAt':'2026-09-11','source':source,'files':files,'profiles':len(profiles),'values':36,'overlapChecks':12,'visualReview':['output/sasol-analysts-2016.pdf-page-20.png','output/sasol-analysts-2019.pdf-page-41.png'],'scope':scope})
ledger=read(HERE/'remaining-40-ledger.json');topic=next(t for t in ledger['topics'] if t['id']=='fuels:synthetic_fuels');topic['research']=[{'sourceId':ID,'reviewedAt':'2026-09-11','evidence':'public-synfuels-audit.json'}]
if topic['status']=='research_pending':topic['status']='source_validated'
write(HERE/'remaining-40-ledger.json',ledger)
print(json.dumps({'source':ID,'profiles':4,'originalFacts':36,'overlappingCellsChecked':12,'bytes':len(package)}))
