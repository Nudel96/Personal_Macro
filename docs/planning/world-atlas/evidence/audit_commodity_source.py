"""Build the reviewed Pink Sheet catalog from original public workbooks.

No personal data. Download inputs separately to the ignored validation directory.
The manual labels, topic links and source-boundary years are editorial metadata.
"""
from collections import Counter
from decimal import Decimal
from hashlib import sha256
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[4]
INPUT = ROOT / "apps/desktop/.tmp/atlas-validation/commodities"
OUTPUT = ROOT / "apps/desktop/src/features/world-atlas/data/commodity-catalog.json"

# One entry for each original price column B:BR, in source order.
PRICES = """
oil_average|Rohöl · Durchschnitt|energy|Crude oil, average
oil_brent|Rohöl · Brent|energy|Crude oil, UK Brent
oil_dubai|Rohöl · Dubai|energy|Crude oil, Dubai
oil_wti|Rohöl · WTI|energy|Crude oil, US,
coal_australia|Kohle · Australien|energy|Coal (Australia)
coal_south_africa|Kohle · Südafrika|energy|Coal (South Africa)
gas_us|Erdgas · USA|energy|Natural Gas (U.S.)
gas_europe|Erdgas · Europa|energy|Natural Gas (Europe)
lng_japan|Flüssigerdgas · japanischer Importpreis|energy|Liquefied natural gas (Japan)
gas_index|Erdgas · veröffentlichter Index|energy|Natural gas index
cocoa|Kakao|beverages|Cocoa (ICCO)
coffee_arabica|Kaffee · Arabica|beverages|Coffee, Arabica
coffee_robusta|Kaffee · Robusta|beverages|Coffee, Robusta
tea_average|Tee · drei Auktionen|beverages|Tea, average
tea_colombo|Tee · Colombo|beverages|Tea (Colombo
tea_kolkata|Tee · Kolkata|beverages|Tea (Kolkata
tea_mombasa|Tee · Mombasa|beverages|Tea (Mombasa
coconut_oil|Kokosöl|oils|Coconut oil,
groundnuts|Erdnüsse|oils|Peanut (Groundnut),
fish_meal|Fischmehl|oils|Fish meal,
groundnut_oil|Erdnussöl|oils|Peanut (Groundnut) oil,
palm_oil|Palmöl|oils|Palm oil (Malaysia)
palmkernel_oil|Palmkernöl|oils|Palmkernel oil (Malaysia)
soybeans|Sojabohnen|oils|Soybeans
soybean_oil|Sojaöl|oils|Soybean oil,
soybean_meal|Sojaschrot|oils|Soybean meal,
barley|Gerste|grains|Barley
maize|Mais|grains|Maize
sorghum|Sorghum|grains|Sorghum
rice_thai_5|Reis · Thailand 5 % Bruchanteil|grains|Rice (Thailand), 5%
rice_thai_25|Reis · Thailand 25 % Bruchanteil|grains|Rice (Thailand), 25%
rice_thai_a1|Reis · Thailand A1|grains|Rice (Thailand), 100% broken,
rice_vietnam|Reis · Vietnam 5 % Bruchanteil|grains|Rice (Vietnam)
wheat_srw|Weizen · USA SRW|grains|Wheat (U.S.), no. 2, soft
wheat_hrw|Weizen · USA HRW|grains|Wheat (U.S.), no. 2 hard
banana_europe|Bananen · Europa|food|Bananas (Central & South America), major brands, free
banana_us|Bananen · USA|food|Bananas (Central & South America), major brands, US
orange|Orangen|food|Orange
beef|Rindfleisch|food|Beef
chicken|Hähnchenfleisch|food|Chicken
lamb|Lammfleisch|food|Lamb
shrimp|Garnelen · US-Referenzpreis|food|Shrimp
sugar_eu|Zucker · EU|food|Sugar (EU)
sugar_us|Zucker · USA|food|Sugar (U.S.)
sugar_world|Zucker · Weltmarkt|food|Sugar (World)
tobacco|Tabak · US-Importpreis|industrial|Tobacco (any origin)
logs_cameroon|Rundholz · Kamerun|timber|Logs (Africa)
logs_malaysia|Rundholz · Malaysia|timber|Logs (Southeast Asia)
sawnwood_cameroon|Schnittholz · Kamerun|timber|Sawnwood (Africa)
sawnwood_malaysia|Schnittholz · Südostasien|timber|Sawnwood (Southeast Asia)
plywood|Sperrholz|timber|Plywood (Africa
cotton|Baumwolle · A-Index|industrial|Cotton (Cotton Outlook
rubber_tsr20|Naturkautschuk · TSR20|industrial|Rubber (Asia), TSR 20
rubber_rss3|Naturkautschuk · RSS3|industrial|Rubber (Asia), RSS3
phosphate|Phosphatgestein|fertilizers|Phosphate rock
dap|Dünger · DAP|fertilizers|DAP (diammonium
tsp|Dünger · TSP|fertilizers|TSP (triple
urea|Harnstoff|fertilizers|Urea,
potash|Kaliumchlorid|fertilizers|Potassium chloride (muriate
aluminum|Aluminium|metals|Aluminum (LME)
iron_ore|Eisenerz|metals|Iron ore (any origin)
copper|Kupfer|metals|Copper (LME)
lead|Blei|metals|Lead (LME)
tin|Zinn|metals|Tin (LME)
nickel|Nickel|metals|Nickel (LME)
zinc|Zink|metals|Zinc (LME)
gold|Gold|precious|Gold,
platinum|Platin|precious|Platinum,
silver|Silber|precious|Silver (UK)
""".strip()
INDEX_LABELS = [
    ("all", "Rohstoffe insgesamt"), ("energy", "Energie"), ("non_energy", "Rohstoffe ohne Energie"),
    ("agriculture", "Agrarrohstoffe"), ("beverages", "Getränkerohstoffe"), ("food", "Nahrungsmittel"),
    ("oils", "Öle und Schrote"), ("grains", "Getreide"), ("other_food", "Weitere Nahrungsmittel"),
    ("raw", "Agrarische Rohmaterialien"), ("timber", "Holz"), ("other_raw", "Weitere Rohmaterialien"),
    ("fertilizers", "Düngemittel"), ("metals", "Metalle und Mineralien"),
    ("base_metals", "Basismetalle ohne Eisenerz"), ("precious", "Edelmetalle"),
]
GROUPS = [
    ("indices", "Breite Rohstoffgruppen"), ("energy", "Energie"), ("metals", "Industriemetalle"),
    ("precious", "Edelmetalle"), ("fertilizers", "Düngemittel"), ("grains", "Getreide"),
    ("oils", "Öle und Schrote"), ("beverages", "Kaffee, Kakao und Tee"), ("food", "Weitere Lebensmittel"),
    ("timber", "Holz"), ("industrial", "Weitere Agrarrohstoffe"),
]

# A change during a year isolates that mixed annual observation on both sides.
BREAKS = {
    "oil_average": [1985], "oil_dubai": [1985], "coal_australia": [2002, 2015, 2022],
    "coal_south_africa": [1990, 2002, 2006, 2015, 2017, 2018], "gas_europe": [2000, 2010, 2015],
    "coconut_oil": [1999, 2021, 2024, 2025], "groundnuts": [1999], "fish_meal": [1999, 2021],
    "groundnut_oil": [1999, 2020, 2023, 2024, 2025], "palm_oil": [2001, 2021, 2024, 2025],
    "palmkernel_oil": [2001, 2021, 2024], "soybean_meal": [1990, 1999, 2021, 2025],
    "soybean_oil": [1999, 2021, 2025], "soybeans": [2007, 2021, 2025], "barley": [1980, 2012],
    "rice_thai_a1": [2006], "wheat_hrw": [2020], "banana_europe": [2006], "shrimp": [2004],
    "chicken": [2021], "beef": [2024], "lamb": [2020],
    "tobacco": [2025], "logs_cameroon": [1996], "logs_malaysia": [1993],
    "sawnwood_malaysia": [2005], "cotton": [2006], "rubber_rss3": [2000, 2004],
    "potash": [2020], "urea": [2022], "aluminum": [2005], "iron_ore": [2008, 2009],
    "nickel": [2005], "zinc": [1990], "gold": [2025], "platinum": [2025], "silver": [1962, 1976],
}
TOPICS = {
    "materials:commodity_prices": {"group": "indices", "ids": []},
    "fuels:oil": {"group": "energy", "ids": ["oil_average", "oil_brent", "oil_dubai", "oil_wti"]},
    "fuels:coal": {"group": "energy", "ids": ["coal_australia", "coal_south_africa"]},
    "fuels:natural_gas": {"group": "energy", "ids": ["gas_europe", "gas_us", "gas_index"]},
    "fuels:lng": {"group": "energy", "ids": ["lng_japan"]},
    "materials:mining": {"group": "metals", "ids": []},
    "materials:precious_metals": {"group": "precious", "ids": []},
    "materials:copper": {"group": "metals", "ids": ["copper"]},
    "materials:aluminium": {"group": "metals", "ids": ["aluminum"]},
    "materials:nickel": {"group": "metals", "ids": ["nickel"]},
    "materials:fertilizers": {"group": "fertilizers", "ids": []},
    "materials:forestry": {"group": "timber", "ids": []},
    "food_water:fisheries": {"group": "food", "ids": ["fish_meal", "shrimp"]},
}


def main():
    raw = (INPUT / "annual.xlsx").read_bytes()
    download = json.loads((INPUT / "download.json").read_text(encoding="utf-8"))
    assert sha256(raw).hexdigest() == download["sha256"]
    book = load_workbook(INPUT / "annual.xlsx", data_only=True, read_only=True)
    sheets = {name: list(book[name].iter_rows(min_row=1, max_row=80, max_col=72)) for name in [
        "Annual Prices (Nominal)", "Annual Prices (Real)", "Annual Indices (Nominal)", "Annual Indices (Real)"
    ]}
    descriptions = []
    for row in book["Description"].iter_rows(values_only=True):
        if len(row) > 19 and isinstance(row[1], str) and isinstance(row[19], str):
            descriptions.append({"text": row[1], "source": row[19]})
    definitions = [row[0] for row in book["Definitions"].iter_rows(values_only=True) if row[0]]
    notes = [" ".join(str(v) for v in row if v is not None) for row in book["Description"].iter_rows(values_only=True)]
    notes = [n for n in notes if n]
    metrics = []
    for column, entry in enumerate(PRICES.splitlines(), 2):
        id_, label, group, prefix = entry.split("|")
        found = [d for d in descriptions if d["text"].startswith(prefix)]
        if len(found) != 1:
            raise ValueError((id_, prefix, [d["text"] for d in found]))
        metrics.append({
            "id": id_, "label": label, "group": group, "kind": "price", "column": column,
            "nominalLabel": sheets["Annual Prices (Nominal)"][6][column-1].value,
            "realLabel": sheets["Annual Prices (Real)"][6][column-1].value,
            "unit": sheets["Annual Prices (Nominal)"][7][column-1].value,
            "description": found[0]["text"], "source": found[0]["source"],
            "boundaryYears": BREAKS.get(id_, []), "firstComparableYear": 2009 if id_ == "iron_ore" else 1960,
        })
    assert len(metrics) == 69
    for column, (id_, label) in enumerate(INDEX_LABELS, 2):
        labels = [sheets["Annual Indices (Nominal)"][r][column-1].value for r in range(5,9)]
        source_label = next(v for v in labels if isinstance(v,str) and v.strip())
        metrics.append({"id": "index_"+id_, "label": label+" · Index", "group": "indices", "kind": "index", "column": column,
            "nominalLabel": source_label, "realLabel": source_label, "unit": "2010=100", "description": definitions[4],
            "source": "World Bank", "boundaryYears": [2018] if id_ in ["all","non_energy","agriculture","raw","other_raw"] else [], "firstComparableYear":1960})
    for selection in TOPICS.values():
        assert all(any(m["id"] == id_ for m in metrics) for id_ in selection["ids"])
    points = []
    counts = Counter()
    for year in range(1960,2026):
        values = {}
        for metric in metrics:
            for basis in ["nominal", "real"]:
                title = f'Annual {"Prices" if metric["kind"] == "price" else "Indices"} ({basis.title()})'
                row = year - 1960 + (8 if metric["kind"] == "price" else 9)
                assert sheets[title][row][0].value == year
                value = sheets[title][row][metric["column"]-1].value
                if value in [None, "…", ".."]:
                    fixed = None
                    counts["missing"] += 1
                else:
                    assert isinstance(value,(int,float))
                    exact = Decimal(str(value))*100
                    assert exact >= 0 and exact == exact.to_integral_value()
                    fixed = int(exact)
                    counts["numeric"] += 1
                values[f'{metric["id"]}:{basis}'] = fixed
        points.append({"year":year,"values":values})
    # Use each cell's published display precision, not a fixed tolerance.
    # Ratios are positive, so interval extrema follow directly from numerator
    # and denominator endpoints. This check never replaces published values.
    def interval(cell):
        assert cell.number_format in ('0', '0.0', '0.00')
        places = len(cell.number_format.split('.')[1]) if '.' in cell.number_format else 0
        half = Decimal(5) * Decimal(10) ** (-places-1)
        value = Decimal(str(cell.value))
        return value-half, value+half
    reconciliation = {'checked':0, 'issues':[]}
    for kind, columns, start, muv_column in [('Prices',69,8,70),('Indices',16,9,17)]:
        nominal_sheet = sheets[f'Annual {kind} (Nominal)']
        real_sheet = sheets[f'Annual {kind} (Real)']
        for year in range(1960,2026):
            row = start + year-1960
            muv = real_sheet[row][muv_column]
            assert muv.value == sheets['Annual Prices (Real)'][8+year-1960][70].value
            ml,mh=interval(muv)
            for column in range(1,columns+1):
                nominal, real = nominal_sheet[row][column], real_sheet[row][column]
                if not isinstance(nominal.value,(int,float)) or not isinstance(real.value,(int,float)): continue
                reconciliation['checked']+=1
                nl,nh=interval(nominal);rl,rh=interval(real)
                if rh < nl*100/mh or rl > nh*100/ml:
                    reconciliation['issues'].append({'family':kind,'year':year,'column':column+1,'nominal':nominal.value,'real':real.value,'muv':muv.value,'difference':float(Decimal(str(real.value))-Decimal(str(nominal.value))*100/Decimal(str(muv.value)))})
    assert reconciliation['checked']==5155
    assert [(i['family'],i['year'],i['column']) for i in reconciliation['issues']]==[('Indices',2014,8)]
    assert all(v is not None and v>0 for v in points[50]['values'].values())
    (INPUT/'real-reconciliation.json').write_text(json.dumps(reconciliation,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    config = {
        "datasetId":"worldbank-commodity-prices", "version":"2026-09-09.1", "recipe":"worldbank-pink-sheet-annual-20260902-v1",
        "release":"September 2026", "sourceRelease":"Updated on September 02, 2026",
        "sourceUrl":"https://www.worldbank.org/en/research/commodity-markets", "url":download["url"],
        "bytes":len(raw), "sha256":sha256(raw).hexdigest(), "firstYear":1960,"lastYear":2025,
        "expectedNumericCells":counts["numeric"],"expectedMissingCells":counts["missing"],
        "groups":[{"id":id_,"label":label} for id_,label in GROUPS],"metrics":metrics,"topics":TOPICS,
        "definitions":definitions,"sourceNotes":notes,
        "sourceDiscrepancies":[{"metric":"index_oils","basis":"real","year":2014,"valueHundredths":9830,"note":"Der veröffentlichte Realindex lässt sich mit Nominalindex und MUV desselben Jahres auch innerhalb ihrer Rundungsintervalle nicht vollständig nachvollziehen. Der Quellenwert bleibt unverändert gekennzeichnet."}],
    }
    OUTPUT.write_text(json.dumps(config,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (INPUT/'reference.json').write_text(json.dumps({"points":points},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    report={"source":download,"firstYear":1960,"lastYear":2025,"metricCount":len(metrics),"numericCells":counts["numeric"],"missingCells":counts["missing"],"fixedPrecision":"exact source values in hundredths; no float rounding", "comparisonBase":"2010 exists for all 85 metrics; reference only, not fair value", "realReconciliation":json.loads((INPUT/'real-reconciliation.json').read_text(encoding='utf-8')), "limits":["International reference quotations, not country sector valuations","Source specification changes and estimates remain visible","Iron ore before 2009 excluded from comparable picture because the source description mixes quotation units; source values retained","Workbook cover mentions comparison sheets not included in the published file; cover not used as numerical metadata"]}
    Path(__file__).with_name('commodity-source-audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print({"metrics":len(metrics),**counts})


if __name__ == '__main__':
    main()
