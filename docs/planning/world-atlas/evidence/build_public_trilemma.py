"""Independently audit the original Aizenman/Chinn/Ito 2020 workbook."""
import hashlib
import json
from collections import defaultdict
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile
import openpyxl

ROOT = Path(__file__).resolve().parents[4]
HERE = Path(__file__).resolve().parent
DATA = ROOT / 'apps/desktop/src/features/world-atlas/data'
RAW = ROOT / 'apps/desktop/.tmp/atlas-remaining-40'
ID = 'aci-trilemma'
FILE = 'aci-trilemma-2020.xlsx'
URL = 'https://web.pdx.edu/~ito/trilemma_indexes_update2020.xlsx'
VERSION = '2026-09-11.19'
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'

def read(p):
    return json.loads(p.read_text(encoding='utf-8'))

def write(p, value):
    p.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

book = openpyxl.load_workbook(RAW / FILE, read_only=True, data_only=True)
rows = list(book.active.values)
assert book.sheetnames == ['Sheet1'] and len(rows) == 12078 and len(rows[0]) == 6
assert rows[0] == ('IMF-World Bank Country Code', 'year', 'Exchange Rate Stability Index',
                   'Monetary Independence Index', 'Financial Openness Index', 'Country Name')
archive = ZipFile(RAW / FILE)
raw_cells = {}
with archive.open('xl/worksheets/sheet1.xml') as stream:
    for _, e in ET.iterparse(stream, events=('end',)):
        if e.tag == NS + 'c':
            address = e.attrib['r']
            v = e.find(NS + 'v')
            if address[0] in 'CDE' and e.attrib.get('t') is None and v is not None:
                assert e.find(NS + 'f') is None and address not in raw_cells
                raw_cells[address] = v.text
            e.clear()
        elif e.tag == NS + 'row':
            e.clear()

# Independent public IMF-code/name/ISO3 crosswalk, audited against the ACI names.
gfdd = openpyxl.load_workbook(RAW / 'worldbank-gfdd-2022.xlsx', read_only=True, data_only=True)
crosswalk = {}
for r in gfdd['Data - August 2022'].iter_rows(min_row=2, values_only=True):
    if r[2] is not None:
        code = str(int(r[2])); value = (r[0], r[3])
        assert crosswalk.setdefault(code, value) == value
aliases = {'516': ('Brunei', 'Brunei Darussalam'), '528': ('Taiwan', 'Taiwan, China'),
           '532': ('Hong Kong, China', 'Hong Kong SAR, China'), '624': ('Cape Verde', 'Cabo Verde'),
           '662': ("C?e d'Ivoire", "Côte d'Ivoire"), '716': ('S? Tom�and Principe', 'São Tomé and Principe'),
           '734': ('Swaziland', 'Eswatini'), '942': ('Serbia, Rep. of', 'Serbia')}
unmapped = {'163': 'Euro_Area', '353': 'Netherlands Antilles', '859': 'American Samoa', '934': 'Czechoslovakia'}
geos = {g['iso3']: g for g in read(DATA / 'catalog.json')['geographies'] if g.get('iso3')}
scope = ('Drei getrennte Forschungsindizes nach Aizenman, Chinn und Ito. Ausgabe 31.08.2021; '
         'Quellenjahre bis 2020, rechtliche Kapitalverkehrsoffenheit nur bis 2019. '
         'Die Skala 0–1 beschreibt jeweils eine Dimension, keine politische Güte, faire Bewertung oder Gesamtregimenote. '
         'Länderbezeichnungen und historische Abgrenzungen folgen der Quelle; keine Rückrechnung auf heutige Grenzen. '
         'USA, Welt und Kontinente besitzen kein eigenes Bild. Keine ergänzten Jahre oder Mischung mit späteren Indexausgaben. '
         'Quelle: Aizenman, Chinn und Ito (2010), The Emerging Global Financial Architecture: Tracing and Evaluating '
         'the New Patterns of the Trilemma’s Configurations, Journal of International Money and Finance 29(4), 615–641.')
definitions = [
    ('ers', 2, 'C', 'Wechselkursstabilität',
     'Wie stabil sich der Wechselkurs gegenüber der Bezugswährung im veröffentlichten Forschungsmodell verhält.',
     'ERS: höhere Werte bedeuten stabilere Wechselkurse zur Bezugswährung. Das Modell verwendet Volatilität und Bindungsschwellen; kein Wechselkursniveau.'),
    ('mi', 3, 'D', 'Geldpolitische Unabhängigkeit',
     'Wie unabhängig sich kurzfristige Zinsen von denen des jeweiligen Bezugslands bewegen.',
     'MI: höher bedeutet größere geldpolitische Unabhängigkeit im Modell. Die Quelle glättet mit Vorjahr, aktuellem Jahr und Folgejahr. Konstante Zinsen können den Modellwert 0,5 erzeugen; er bedeutet keine politische Neutralität. Kein in Echtzeit verfügbares Handelssignal.'),
    ('kaopen', 4, 'E', 'Rechtliche Kapitalverkehrsoffenheit',
     'Wie offen die veröffentlichten rechtlichen Regeln für grenzüberschreitenden Kapitalverkehr sind.',
     'KAOPEN: höher bedeutet rechtlich offeneren Kapitalverkehr. Modellindex aus gemeldeten Vorschriften, keine beobachteten Kapitalströme. Diese gemeinsame Ausgabe endet hier 2019; sie wird nicht mit der späteren eigenständigen Chinn-Ito-Ausgabe vermischt.'),
]
# Deliberate display breaks at major changes of territory; never provider flags.
boundaries = {'134': [1991], '513': [1972], '564': [1972], '643': [1993], '644': [1993],
              '728': [1990], '732': [2011, 2012], '474': [1990]}
boundary_note = 'Historische Gebietsabgrenzung der Quelle; vor und nach Gebietsänderungen keine Rekonstruktion identischer heutiger Landesgrenzen. Atlas trennt bekannte Übergänge in der Linie.'
metrics, columns = [], []
for code, index, col, label, explanation, note in definitions:
    metrics.append({'id': ID + ':' + code, 'sourceId': ID, 'topicId': 'institutions:economic_policy_regimes',
                    'providerCode': code, 'label': label + ' · ACI', 'unit': 'Forschungsindex (0–1)',
                    'frequency': 'annual', 'kind': 'model_estimate', 'comparison': 'same_definition',
                    'connectAdjacent': True, 'explanation': explanation, 'scopeNote': note + ' ' + scope})
    columns.append({'code': code, 'index': index, 'column': col, 'title': rows[0][index], 'note': note})
identities, seen, values = {}, set(), defaultdict(dict)
for row_num, r in enumerate(rows[1:], 2):
    code, year, name = str(int(r[0])), r[1], r[5]
    assert isinstance(year, int) and 1960 <= year <= 2020 and (code, year) not in seen
    seen.add((code, year))
    assert identities.setdefault(code, name) == name
    if code in unmapped:
        assert name == unmapped[code]
    else:
        iso, reference_name = crosswalk[code]
        assert name == reference_name or aliases[code] == (name, reference_name)
        assert iso in geos
    for col in columns:
        raw = raw_cells.get(col['column'] + str(row_num))
        value = r[col['index']]
        assert (raw is None) == (value is None)
        if raw is not None:
            assert 0 <= Decimal(raw) <= 1 and float(raw) == value
        values[(code, col['code'])][year] = raw
assert len(identities) == 199 and len(raw_cells) == 24761
profiles, areas, countries = [], {}, {}
for code, name in identities.items():
    iso = crosswalk[code][0] if code in crosswalk else None
    countries[code] = {'label': name, 'iso3': iso, 'geographyId': geos[iso]['id'] if iso else None,
                       'breakYears': boundaries.get(code, []), 'note': boundary_note if code in boundaries else ''}
for (code, metric_code), series in sorted(values.items()):
    if code in unmapped or all(v is None for v in series.values()):
        continue
    c = countries[code]
    col = next(x for x in columns if x['code'] == metric_code)
    metric = next(m for m in metrics if m['providerCode'] == metric_code)
    a = areas.setdefault(code, {'code': code, 'label': c['label'], 'geographyId': c['geographyId'], 'seriesTitles': {}})
    a['seriesTitles'][metric_code] = col['title']
    points = [{'period': str(y), 'value': v, 'status': 'Veröffentlichter Forschungsindex' if v is not None else 'Kein veröffentlichter Zahlenwert',
               'breakBefore': y in c['breakYears'], 'notes': [col['note']] + ([c['note']] if c['note'] else []),
               'lowerBound': None, 'upperBound': None} for y, v in sorted(series.items())]
    profiles.append({'metricId': metric['id'], 'geographyId': a['geographyId'], 'providerArea': code,
                     'providerLabel': a['label'], 'providerTitle': col['title'], 'unit': metric['unit'], 'points': points})
profiles.sort(key=lambda p: (p['metricId'], p['geographyId']))
numeric = sum(p['value'] is not None for a in profiles for p in a['points'])
source = {'id': ID, 'label': 'Aizenman · Chinn · Ito · Trilemma-Indizes', 'adapter': 'aci_trilemma',
          'url': URL, 'documentationUrl': 'https://web.pdx.edu/~ito/trilemma_indexes.htm',
          'licenseUrl': 'https://web.pdx.edu/~ito/ReadMe_trilemma_indexes2014.pdf',
          'publishedAt': '2021-08-31', 'reviewedAt': '2026-09-11', 'recipe': 'aci-2020-original-three-dimensions-v1',
          'observationKind': 'published_research_indices', 'expectedSha256': hashlib.sha256((RAW / FILE).read_bytes()).hexdigest(),
          'expectedRows': 12077, 'expectedNumeric': numeric, 'firstPeriod': '1960', 'lastPeriod': '2020', 'areas': list(areas.values())}
contract = {'sourceId': ID, 'file': FILE, 'sheet': 'Sheet1', 'header': rows[0], 'columns': columns,
            'countries': countries, 'originalNumeric': 24761, 'zipEntries': len(archive.infolist()),
            'uncompressedBytes': sum(a.file_size for a in archive.infolist())}
catalog = read(DATA / 'public-series-catalog.json')
catalog['sources'] = [s for s in catalog['sources'] if s['id'] != ID] + [source]
catalog['metrics'] = [m for m in catalog['metrics'] if m['sourceId'] != ID] + metrics
catalog['version'] = VERSION; write(DATA / 'public-series-catalog.json', catalog)
main = read(DATA / 'catalog.json'); main['version'] = VERSION; write(DATA / 'catalog.json', main)
write(DATA / 'public-trilemma-contract.json', contract)
write(RAW / 'trilemma-expected-profiles.json', profiles)
audit = {'reviewedAt': '2026-09-11', 'source': source, 'contract': contract, 'profiles': len(profiles),
         'numeric': numeric, 'missing': sum(p['value'] is None for a in profiles for p in a['points']),
         'originalDecimalsPreserved': True, 'crosswalkOriginal': 'GFDD original 20220909, imfn/iso3/country fields; individually reviewed aliases',
         'excluded': unmapped, 'boundary': scope}
write(HERE / 'public-trilemma-audit.json', audit)
ledger = read(HERE / 'remaining-40-ledger.json')
topic = next(t for t in ledger['topics'] if t['id'] == 'institutions:economic_policy_regimes')
topic['research'] = [{'sourceId': ID, 'reviewedAt': '2026-09-11', 'evidence': 'public-trilemma-audit.json'}]
if topic['status'] == 'research_pending': topic['status'] = 'source_validated'
write(HERE / 'remaining-40-ledger.json', ledger)
print(json.dumps({'source': ID, 'areas': len(areas), 'profiles': len(profiles), 'numeric': numeric, 'missing': audit['missing']}))
