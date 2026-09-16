"""Audit the published OECD HM1.1.A1 worksheet, never its old working sheets."""
import argparse
import hashlib
import json
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile
import openpyxl

ROOT = Path(__file__).resolve().parents[4]
TMP = ROOT / 'apps/desktop/.tmp/atlas-remaining-40'
DATA = ROOT / 'apps/desktop/src/features/world-atlas/data'
EVIDENCE = Path(__file__).resolve().parent
URL = 'https://webfs.oecd.org/Els-com/Affordable_Housing_Database/HM1-1-Housing-stock-and-construction.xlsx'
DOC = URL.replace('.xlsx', '.pdf')
SHEET = 'HM1.1.A1'
SOURCE = 'oecd-housing-stock'
COUNTRIES = dict(zip(
    ['Australia','Austria','Belgium','Brazil','Bulgaria','Canada','Chile','Colombia','Costa Rica','Croatia','Cyprus','Czechia','Denmark','Estonia','Finland','France','Germany','Greece','Hungary','Iceland','Ireland','Italy','Japan','Korea','Latvia','Lithuania','Luxembourg','Malta','Netherlands','New Zealand','Norway','Poland','Portugal','Romania','Slovak Republic','Slovenia','South Africa','Spain','Sweden','Switzerland','Türkiye','UK (England)','United States'],
    ['AUS','AUT','BEL','BRA','BGR','CAN','CHL','COL','CRI','HRV','CYP','CZE','DNK','EST','FIN','FRA','DEU','GRC','HUN','ISL','IRL','ITA','JPN','KOR','LVA','LTU','LUX','MLT','NLD','NZL','NOR','POL','PRT','ROU','SVK','SVN','ZAF','ESP','SWE','CHE','TUR','ENGLAND','USA'], strict=True))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()
    path = TMP / 'oecd-housing-stock.xlsx'
    raw = path.read_bytes()
    book = openpyxl.load_workbook(path, data_only=True)
    sheet = book[SHEET]
    assert sheet['A1'].value == 'Table HM1.1.A1. Total housing stock in OECD and EU countries, selected years'
    assert sheet['B4'].value == 'Total number of dwellings'
    assert sheet['H4'].value == 'Dwellings per 1000 inhabitants'
    # Independently inspect the XML cells backing openpyxl's cached values.
    ns = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(path) as z:
        workbook = ET.fromstring(z.read('xl/workbook.xml'))
        item = next(s for s in workbook.find('s:sheets',ns) if s.attrib['name'] == SHEET)
        rels = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        rid = item.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
        target = next(r.attrib['Target'] for r in rels if r.attrib['Id'] == rid)
        xml = ET.fromstring(z.read('xl/' + target))
        cells = {c.attrib['r']: c for c in xml.findall('.//s:c',ns)}
    catalog = json.loads((DATA/'catalog.json').read_text(encoding='utf-8'))
    geographies = {g['iso3']:g for g in catalog['geographies'] if g['iso3']}
    england = dict(id='oecd:england', label='England · OECD-Wohnungsstatistik', iso3='', regionId='Europe', kind='aggregate')
    metrics = []
    for code, label, unit, explanation in [
        ('total', 'Wohnungsbestand · insgesamt', 'Wohnungen', 'Veröffentlichter Gesamtbestand an Wohnungen, einschließlich bewohnter und unbewohnter Einheiten. Das ist kein aktuelles Angebot an freien Miet- oder Kaufobjekten.'),
        ('per_1000', 'Wohnungsbestand je 1.000 Einwohner', 'Wohnungen je 1.000 Einwohner', 'Von der OECD veröffentlichte Relation des Wohnungsbestands zur Bevölkerung. Die Quote wird unverändert übernommen; sie ist keine Leerstands- oder Eigentumsquote.'),
    ]:
        metrics.append(dict(id=SOURCE+':'+code, sourceId=SOURCE, topicId='housing:housing_supply', providerCode=code,
                            label=label, unit=unit, frequency='annual', kind='numeric', comparison='same_definition', connectAdjacent=False,
                            explanation=explanation, scopeNote='OECD Affordable Housing Database, HM1.1.A1, Ausgabe vom 15. April 2024. Erhebungen und nationale Schätzungen zu ausgewählten Jahren; die tatsächlichen Bezugsjahre gelten statt der Tabellenüberschriften „um 2011“, „um 2018“ oder „2022“. Einzelpunkte bleiben unverbunden. Bestandsdefinition und Erfassung können sich zwischen Ländern unterscheiden. England bleibt vom Vereinigten Königreich getrennt; Zypern erfasst das von seiner Regierung kontrollierte Gebiet. Keine Markt- oder Immobilienbewertung.'))
    specs, areas, profiles, audit_rows = [], [], [], []
    undated = 0
    for row, (label, iso) in enumerate(COUNTRIES.items(), 6):
        assert sheet.cell(row,1).value == label
        geography = england if iso == 'ENGLAND' else geographies[iso]
        area = dict(code=label, label=label, geographyId=geography['id'], seriesTitles={'total': 'Total number of dwellings', 'per_1000':'Dwellings per 1000 inhabitants'})
        areas.append(area)
        years = [sheet.cell(row,c).value for c in [11,12,13]]
        assert len([y for y in years if y is not None]) == len(set(y for y in years if y is not None))
        specs.append(dict(row=row-1, label=label, geographyId=geography['id'], years=years))
        for metric, offset in zip(metrics,[2,8]):
            points=[]
            for index, year in enumerate(years):
                value=sheet.cell(row,offset+index).value
                if year is None:
                    assert value is None
                    undated += 1
                    continue
                assert isinstance(year,int) and 2010 <= year <= 2022
                assert isinstance(value,(int,float)) and value > 0
                if metric['providerCode']=='total': assert value == int(value)
                cell = sheet.cell(row,offset+index).coordinate
                original = cells[cell].find('s:v',ns).text
                assert float(original) == value
                notes = [f'OECD HM1.1.A1 · tatsächliches Bezugsjahr {year}.', 'Originaltabellen und nationale Erhebungen bzw. Schätzungen; einzelne Quellenjahre.']
                if iso=='ENGLAND': notes.append('Quellengebiet: nur England, nicht das gesamte Vereinigte Königreich.')
                if iso=='CYP': notes.append('Quellengebiet: von der Regierung der Republik Zypern kontrolliertes Gebiet.')
                if iso=='COL': notes.append('Kolumbien: Bestandsprojektionen von DANE sind in dieser Quelle enthalten.')
                if iso=='ESP': notes.append('Spanien: nationale Schätzung des Wohnungsbestands.')
                # Retain both published table notes, including their source/year qualifications.
                notes.extend([sheet['A50'].value, sheet['A51'].value])
                points.append(dict(period=str(year), value=str(value), status='', breakBefore=False, notes=notes, lowerBound=None, upperBound=None))
                audit_rows.append(dict(geographyId=geography['id'],metric=metric['providerCode'],cell=cell,year=year,xmlValue=original))
            profiles.append(dict(metricId=metric['id'], geographyId=geography['id'], providerArea=label, providerLabel=label,
                                 providerTitle=area['seriesTitles'][metric['providerCode']],unit=metric['unit'],points=points))
    contract = dict(sourceId=SOURCE, sheet=SHEET, rawFile=path.name,
                    headers=[dict(row=sheet[cell].row-1,column=sheet[cell].column-1,text=sheet[cell].value) for cell in ['A1','A2','B4','E4','H4','K4','B5','C5','D5','H5','I5','J5','K5','L5','M5']],
                    notes=[dict(row=49,text=sheet['A50'].value),dict(row=50,text=sheet['A51'].value)], areas=specs, expectedUndatedBlocks=undated)
    count=sum(len(p['points']) for p in profiles)
    source=dict(id=SOURCE,label='OECD · Wohnungsbestand',adapter='oecd_housing_stock',url=URL,documentationUrl=DOC,
                licenseUrl='https://www.oecd.org/en/about/terms-conditions.html',publishedAt='2024-04-15',reviewedAt='2026-09-11',
                recipe='oecd-hm11-a1-original-years-20240415-v1',observationKind='source_estimates',
                expectedSha256=hashlib.sha256(raw).hexdigest(),expectedRows=43,expectedNumeric=count,firstPeriod='2010',lastPeriod='2022',areas=areas)
    (TMP/'housing-stock-expected-profiles.json').write_text(json.dumps(profiles,ensure_ascii=False,indent=2),encoding='utf-8')
    (TMP/'housing-stock-contract-draft.json').write_text(json.dumps(contract,ensure_ascii=False,indent=2),encoding='utf-8')
    audit=dict(sourceId=SOURCE,source=source,checkedAt='2026-09-11',url=URL,documentationUrl=DOC,sha256=source['expectedSha256'],numericValues=count,profiles=len(profiles),areaCount=len(areas),
               undatedBlocks=undated,originalWorksheet=SHEET,xmlCellCrosscheck=True,rows=audit_rows,
               limits=['Only the published HM1.1.A1 table, no old working-sheet observations.','England is a separate source geography; no UK substitution.','Actual years; sparse observations; no interpolation or own population denominator.'])
    (EVIDENCE/'public-housing-stock-audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    if args.apply:
        public=json.loads((DATA/'public-series-catalog.json').read_text(encoding='utf-8'))
        public['sources']=[s for s in public['sources'] if s['id']!=SOURCE]+[source]
        public['metrics']=[m for m in public['metrics'] if m['sourceId']!=SOURCE]+metrics
        public['version']='2026-09-11.9'
        (DATA/'public-series-catalog.json').write_text(json.dumps(public,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
        (DATA/'public-housing-stock-contract.json').write_text(json.dumps(contract,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
        catalog['geographies']=[g for g in catalog['geographies'] if g['id']!=england['id']]+[england]
        catalog['version']='2026-09-11.9'
        (DATA/'catalog.json').write_text(json.dumps(catalog,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(f'OECD housing stock: {len(areas)} areas / {len(profiles)} profiles / {count} observations; {undated} undated empty blocks omitted; apply={args.apply}')


if __name__=='__main__':main()
