"""Reproducible source contracts and independent expected values for the gap fill.

Only public source files are read. Expected observations stay in .tmp, not git.
Run after the public downloads documented in GAP-EXPANSION.md.
"""
import collections
import csv
import hashlib
import io
import json
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[4]
HERE = Path(__file__).resolve().parent
DATA = ROOT / "apps/desktop/src/features/world-atlas/data"
RAW = ROOT / "apps/desktop/.tmp/atlas-gaps"
EXPECTED = RAW / "expected"
EXPECTED.mkdir(exist_ok=True)
CATALOG = json.loads((DATA / "catalog.json").read_text("utf-8"))
GEOS = {g["iso3"]: g for g in CATALOG["geographies"] if g["iso3"]}
CONFIG = json.loads((DATA / "public-series-catalog.json").read_text("utf-8"))
CONTRACTS, SOURCES, METRICS, AUDIT = {}, [], [], []

def write(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", "utf-8")

def metric(sid, code, topic, label, unit, explanation, scope, connect=False, comparison="same_definition"):
    result = dict(id=f"{sid}:{code}", sourceId=sid, topicId=topic, providerCode=code, label=label, unit=unit, frequency="annual", kind="source_statistic", comparison=comparison, connectAdjacent=connect, explanation=explanation, scopeNote=scope)
    METRICS.append(result)
    return result

def area(code, label, geo, titles):
    return dict(code=code, label=label, geographyId=geo, seriesTitles=titles)

def point(period, value, status="source_statistics", notes=None, br=False):
    return dict(period=str(period), value=value, status=status, breakBefore=br, notes=notes or [], lowerBound=None, upperBound=None)

def profile(m, a, points):
    return dict(metricId=m["id"], geographyId=a["geographyId"], providerArea=a["code"], providerLabel=a["label"], providerTitle=a["seriesTitles"][m["providerCode"]], unit=m["unit"], points=points)

def source(sid, label, adapter, filename, url, docs, license_url, published, kind, areas, profiles, source_rows):
    periods = [p["period"] for row in profiles for p in row["points"]]
    numeric = sum(p["value"] is not None for row in profiles for p in row["points"])
    result = dict(id=sid, label=label, adapter=adapter, url=url, documentationUrl=docs, licenseUrl=license_url, publishedAt=published, reviewedAt="2026-09-15", recipe=f"{sid}-reviewed-20260915-v1", observationKind=kind, expectedSha256=hashlib.sha256((RAW/filename).read_bytes()).hexdigest(), expectedRows=source_rows, expectedNumeric=numeric, firstPeriod=min(periods), lastPeriod=max(periods), areas=areas)
    assert len({a["code"] for a in areas}) == len(areas)
    assert len({a["geographyId"] for a in areas}) == len(areas)
    assert len({(r["metricId"],r["geographyId"]) for r in profiles}) == len(profiles)
    SOURCES.append(result)
    write(EXPECTED/f"{sid}.json", profiles)
    geo_set = {r["geographyId"] for r in profiles if any(p["value"] is not None for p in r["points"])}
    AUDIT.append(dict(sourceId=sid, sourceFile=filename, sha256=result["expectedSha256"], rows=source_rows, numeric=numeric, profiles=len(profiles), areasWithValues=len(geo_set), africanCountriesWithValues=len([g for g in GEOS.values() if g["id"] in geo_set and g["regionId"]=="Africa"]), firstPeriod=min(periods), lastPeriod=max(periods)))

def imf():
    metadata = json.loads((RAW/"imf-indicators.json").read_text("utf-8"))["indicators"]
    countries = json.loads((RAW/"imf-countries.json").read_text("utf-8"))["countries"]
    selections = [
        ("PVD_LS", "finance:credit_growth", "Private Schulden · Kredite und Schuldpapiere"),
        ("HH_LS", "finance:household_debt", "Haushalte · Kredite und Schuldpapiere"),
        ("NFC_LS", "finance:corporate_debt", "Unternehmen · Kredite und Schuldpapiere"),
        ("Privatedebt_all", "finance:credit_growth", "Private Schulden · alle Instrumente"),
        ("HH_ALL", "finance:household_debt", "Haushalte · alle Schuldinstrumente"),
        ("NFC_ALL", "finance:corporate_debt", "Unternehmen · alle Schuldinstrumente"),
    ]
    for code, topic, label in selections:
        sid = "imf-gdd-" + code.lower()
        meta = metadata[code]
        assert meta["dataset"] == "GDD" and meta["source"] == "Global Debt Database (Sep 2025)" and meta["unit"] == "Percent of GDP"
        raw = json.loads((RAW/f"imf-{code}.json").read_text("utf-8"), parse_float=str, parse_int=str)
        vals = raw["values"][code]
        scope = "IWF Global Debt Database, Ausgabe September 2025, historische Jahreswerte bis 2024. Bruttobestand des nichtfinanziellen privaten Sektors relativ zum BIP; keine jährliche Kreditwachstumsrate, kein Schuldendienst und kein fairer Marktwert. Länder unterscheiden sich in Quellen und Instrumentabdeckung; historische Schätzungen und Quellenwechsel sind möglich. Die API liefert keine datierten Bruchkennzeichen. Alle Instrumente und Kredite/Schuldpapiere sind getrennte Definitionen. Keine Fortsetzung der BIS-Quartalsreihe und keine eigenen Weltmittel."
        explanation = ("Veröffentlichter Bruttoschuldenbestand der privaten nichtfinanziellen Haushalte und Unternehmen" if code in {"PVD_LS", "Privatedebt_all"} else "Veröffentlichter Bruttoschuldenbestand privater Haushalte" if code.startswith("HH") else "Veröffentlichter Bruttoschuldenbestand nichtfinanzieller Unternehmen") + " relativ zur jährlichen Wirtschaftsleistung."
        m = metric(sid, code, topic, label, "% des BIP", explanation, scope)
        areas, profiles = [], []
        for iso, years in sorted(vals.items()):
            assert iso in GEOS and iso in countries
            a = area(iso, countries[iso]["label"], GEOS[iso]["id"], {code: meta["label"]})
            points = []
            for year, val in sorted(years.items()):
                assert 1950 <= int(year) <= 2024
                assert val is None or float(val) >= 0
                points.append(point(year, val, "source_statistics_with_estimates"))
            areas.append(a); profiles.append(profile(m,a,points))
        source(sid, "IWF · " + label, "imf_gdd", f"imf-{code}.json", f"https://www.imf.org/external/datamapper/api/v1/{code}", "https://data.imf.org/Datasets/FAD_GDD", "https://www.imf.org/external/terms.htm", "2025-09-16", "source_statistics", areas, profiles, sum(len(v) for v in vals.values()))
        CONTRACTS[sid] = dict(code=code, metadata=meta, areas=areas)

def wage():
    sid = "ilo-real-wage-growth"
    w = load_workbook(RAW/"ilo-real-wages.xlsx", read_only=True, data_only=True)
    rows = list(w["Real wage growth"].values)
    header = list(rows[0])
    assert header[11:36] == list(range(2000,2025))
    with zipfile.ZipFile(RAW/"ilo-real-wages.xlsx") as z:
        root = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
    ns = {"s":"http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    original = {c.attrib["r"]: c.find("s:v",ns).text for c in root.findall(".//s:c",ns) if c.attrib.get("t") is None and c.find("s:v",ns) is not None}
    wage_topic = next(m["topicId"] for m in CONFIG["metrics"] if m["sourceId"]=="oecd-wages")
    note = "ILO Global Wage Report 2024–25: reale Veränderung durchschnittlicher Bruttolöhne, aus nationalen Quellen und ILO-Schätzungen. Kein Lohnniveau in Dollar oder Kaufkraftstandard. Länder, Arbeitnehmerkreis und Erhebungsgrundlage unterscheiden sich. Auch wiederholte konstante Raten können auf Quellenmodellierung beruhen. Fehlwerte bleiben leer. Die im November 2024 veröffentlichten Teiljahreswerte für 2024 sind ausgeschlossen; historische Grenze 2023."
    m = metric(sid, "real_growth", wage_topic, "Reallohnveränderung · ILO", "% gegenüber Vorjahr", "Veröffentlichte inflationsbereinigte Veränderung des durchschnittlichen Lohnes gegenüber dem Vorjahr.", note)
    areas, profiles, row_map = [], [], []
    from openpyxl.utils import get_column_letter
    for index, r in enumerate(rows[1:],2):
        iso, name = r[5], r[4]
        assert iso in GEOS
        a = area(iso,name,GEOS[iso]["id"],{"real_growth":"Real wage growth"})
        points = []
        for col, year in enumerate(header[11:35],12):
            value = original.get(f"{get_column_letter(col)}{index}")
            assert (value is None) == (r[col-1] is None or r[col-1]=="")
            if value is not None: assert float(value) == r[col-1] and float(value) > -100
            points.append(point(year,value,"estimated"))
        areas.append(a); profiles.append(profile(m,a,points)); row_map.append(dict(row=index, code=iso, label=name))
    source(sid,"ILO · Reallohnveränderung","ilo_wage_xlsx","ilo-real-wages.xlsx","https://www.ilo.org/media/627781/download","https://www.ilo.org/publications/flagship-reports/global-wage-report-2024-25-wage-inequality-decreasing-globally","https://www.ilo.org/global/copyright/lang--en/index.htm","2024-11-28","modeled_estimate",areas,profiles,len(rows)-1)
    CONTRACTS[sid] = dict(sheetNames=w.sheetnames, header=header[:36], rows=row_map, firstYear=2000, lastYear=2023)

def hours():
    sid="ilo-weekly-hours"
    rows=list(csv.DictReader((RAW/"ilo-weekly-hours.csv").read_text("utf-8-sig").splitlines()))
    all_areas={r["ref_area"]:r["ref_area.label"] for r in rows}
    mapping={code:GEOS[code]["id"] for code in all_areas if code in GEOS}
    mapping["KOS"]=GEOS["XKX"]["id"]
    note="ILOSTAT COND: tatsächliche Wochenstunden aller Erwerbstätigen, beide Geschlechter, Originalangaben der bevorzugten nationalen Quelle. Kein OECD-Jahresstundenmaß und keine Vollzeit-Norm. Bezugswoche, Erhebung, Haupt-/Nebenjobs und Gebietsabdeckung können abweichen; die Originalhinweise stehen an jedem Punkt. Erhebungen werden nicht verbunden. Historische Grenze 2024; globale und regionale ILO-Modellwerte sind nicht Teil dieser nationalen Reihe."
    m=metric(sid,"actual_weekly","labor:working_hours","Tatsächliche Wochenarbeitszeit · ILOSTAT","Stunden je Erwerbstätigem und Woche","Durchschnittlich tatsächlich gearbeitete Wochenstunden je Erwerbstätigem einschließlich Selbstständiger.",note,comparison="within_country")
    groups=collections.defaultdict(list)
    seen=set()
    for r in rows:
        if r["ref_area"] not in mapping: continue
        assert "Modelled" not in r["source.label"]
        key=(r["ref_area"],r["time"]);assert key not in seen;seen.add(key)
        assert r["sex"]=="SEX_T" and r["sex.label"]=="Total" and r["indicator"]=="HOW_TEMP_SEX_NB"
        assert 1976<=int(r["time"])<=2024
        assert r["obs_status"] in {"","B"}
        notes=[f"{r['source.label']} ({r['source']})"]+[r[k] for k in ["note_indicator.label","note_source.label"] if r[k]]
        if r["obs_status"]: notes.append(r["obs_status.label"])
        groups[r["ref_area"]].append(point(r["time"],r["obs_value"] or None,"source_statistics",notes,r["obs_status"]=="B"))
    title=rows[0]["indicator.label"]
    areas=[area(c,all_areas[c],mapping[c],{"actual_weekly":title}) for c in sorted(groups)]
    profiles=[profile(m,a,sorted(groups[a["code"]],key=lambda p:p["period"])) for a in areas]
    source(sid,"ILOSTAT · Wochenarbeitszeit","ilo_hours_csv","ilo-weekly-hours.csv","https://rplumber.ilo.org/data/indicator?id=HOW_TEMP_SEX_NB_A&sex=SEX_T&format=.csv&type=both&timefrom=1976&timeto=2024","https://ilostat.ilo.org/methods/concepts-and-definitions/description-wages-and-working-time-statistics/","https://www.ilo.org/global/copyright/lang--en/index.htm","2026-09-09","source_statistics",areas,profiles,len(rows))
    CONTRACTS[sid]=dict(headers=list(rows[0]), allAreas=all_areas, excludedAreas={k:v for k,v in all_areas.items() if k not in mapping}, indicator=title)

def icp():
    sid="worldbank-icp-housing"
    raw=json.loads((RAW/"icp-housing.json").read_text("utf-8"),parse_float=str,parse_int=str)
    rows=raw["source"]["data"]
    assert raw["pages"]=="1" and raw["lastupdated"]=="2024-08-04"
    all_areas={next(v["id"] for v in r["variable"] if v["concept"]=="Country"):next(v["value"] for v in r["variable"] if v["concept"]=="Country") for r in rows}
    # Original ICP benchmark aggregates never become UN/WDI regions.
    extra=json.loads((DATA/"public-geographies.json").read_text("utf-8"))
    mapping={c:GEOS[c]["id"] for c in all_areas if c in GEOS}
    mapping["RUT"]=GEOS["RUS"]["id"]
    for code, label in all_areas.items():
        if code in mapping: continue
        ident="icp:"+code
        if not any(g["id"]==ident for g in extra):
            extra.append(dict(id=ident,label=label+" · ICP",iso3="",regionId="Unassigned",kind="provider_area" if code=="BON" else "aggregate"))
        mapping[code]=ident
    write(DATA/"public-geographies.json",extra)
    note="ICP 2021, revidierte Benchmarks 2017 und 2021. Breite Wohnleistungen einschließlich tatsächlicher und unterstellter Eigentümermieten, Wasser, Strom, Gas und weiterer Brennstoffe; auch individuell zurechenbare öffentliche Leistungen. Kein isolierter Mietpreis, kein Euro/m² und keine Immobilienbewertung. Nationale Ausgaben- und Preisstatistiken mit Schätzungen. Nur dieselben Benchmarkjahre vergleichen; die Weltreferenz je Jahr ist keine Zeitreihe der Mietinflation. Zwischen den Erhebungen wird nicht interpoliert."
    choices=[("PX.WL","housing:rents","Wohnen und Versorgung · Preisniveau","Index (Welt = 100 im Benchmarkjahr)"),("AICZS","housing:affordability","Wohnen und Versorgung · Konsumanteil","% des tatsächlichen Individualkonsums"),("ZS","housing:affordability","Wohnen und Versorgung · BIP-Anteil","% des BIP"),("PCAP.PP","housing:affordability","Wohnen und Versorgung · Pro-Kopf-Volumen","Internationale Dollar je Person (Benchmarkjahr)")]
    ms={c:metric(sid,c,t,l,u,"Originaler ICP-Benchmark für die Ausgabenkomponente Wohnen, Wasser, Strom, Gas und weitere Brennstoffe.",note) for c,t,l,u in choices}
    titles={};groups=collections.defaultdict(list)
    for r in rows:
        v={x["concept"]:x for x in r["variable"]};c=v["Country"]["id"];code=v["Classification"]["id"];year=v["Time"]["value"]
        assert v["Series"]["id"]=="9060000" and year in {"2017","2021"}
        titles[code]=v["Classification"]["value"]
        groups[(c,code)].append(point(year,r["value"],"estimated"))
    areas=[area(c,l,mapping[c],titles) for c,l in sorted(all_areas.items())]
    profiles=[profile(ms[code],a,sorted(groups[(a["code"],code)],key=lambda p:p["period"])) for a in areas for code in ms]
    url="https://api.worldbank.org/v2/sources/90/country/all/series/9060000/classification/PX.WL;AICZS;ZS;PCAP.PP/time/YR2017;YR2021?format=json&per_page=20000"
    source(sid,"Weltbank ICP · Wohnkosten und Versorgung","icp_housing","icp-housing.json",url,"https://www.worldbank.org/en/programs/icp/brief/ICP2021_Methodology_PPP","https://www.worldbank.org/en/about/legal/terms-of-use-for-datasets","2024-08-04","modeled_estimate",areas,profiles,len(rows))
    CONTRACTS[sid]=dict(seriesId="9060000",seriesTitle="9060000:ACTUAL HOUSING, WATER, ELECTRICITY, GAS AND OTHER FUELS",lastUpdated="2024-08-04",classifications=titles)

def wto_fingerprint(rows):
    """Full field content, order independent; duplicates still contribute a hash."""
    assert rows and all('\0' not in v for r in rows for v in r.values())
    header='\0'.join(rows[0]).encode('utf-8')
    hashes=sorted(hashlib.sha256('\0'.join(r.values()).encode('utf-8')).digest() for r in rows)
    return hashlib.sha256(header+b'\0'+b''.join(hashes)).hexdigest()

def wto():
    sid="wto-merchandise"
    with zipfile.ZipFile(RAW/"wto-values.zip") as z:
        assert z.namelist()==["merchandise_values_annual_dataset.csv"]
        rows=list(csv.DictReader(io.StringIO(z.read(z.namelist()[0]).decode("cp1252"),newline="")))
    products={r["ProductCode"]:r["Product"] for r in rows}
    labels={"TO":"Waren insgesamt","AG":"Landwirtschaftliche Waren","MA":"Verarbeitete Waren","MI":"Brennstoffe und Bergbauprodukte","AGFO":"Nahrungsmittel","MACL":"Bekleidung","MATE":"Textilien","MAMTOF":"Büro- und Telekommunikationsgeräte","MACH":"Chemische Erzeugnisse","MAMT":"Maschinen und Transportmittel","MAMTAU":"Automobilprodukte","MAIS":"Eisen und Stahl","MIFU":"Brennstoffe","MAMTTE":"Transportmittel","MAMTOTTL":"Telekommunikationsgeräte","MAMTOTEP":"Datenverarbeitungs- und Bürogeräte","MACHPH":"Pharmazeutische Erzeugnisse","MAMTOTIC":"Integrierte Schaltungen und Elektronikkomponenten"}
    assert set(products)==set(labels)
    note="WTO-Warenhandelsstatistik: Bruttohandelswert mit der Welt, laufende Millionen US-Dollar. Exporte in der Regel FOB, Importe CIF; nationale Abweichungen und WTO-Schätzungen sind möglich. Keine preisbereinigten Mengen, Wertschöpfungsanteile oder direkte bilaterale Lieferabhängigkeit. SITC-Rev.-3-Gruppen überlappen hierarchisch; keine Summe über alle Gruppen bilden. Länder- und Methodenstände können historisch wechseln. Originale Qualitätskennzeichen bleiben sichtbar. Randjahr 2025 enthält je Land und Produkt unterschiedlich vollständige Angaben."
    metrics={}
    for flow in ["AX","AM"]:
        for code in ["TO"]+sorted(k for k in products if k!="TO"):
            pc=f"ITS_MTV_{flow}:{code}"
            title=("Exporte" if flow=="AX" else "Importe")+" · "+labels[code]
            topic=("trade:exports" if flow=="AX" else "trade:imports") if code=="TO" else "trade:supply_chains"
            metrics[pc]=metric(sid,pc,topic,title,"Mio. laufende US-Dollar","Veröffentlichter grenzüberschreitender Warenhandel der gewählten Produktgruppe mit der Welt.",note)
    reporters={r["ReporterCode"]:[r["ReporterISO3A"],r["Reporter"]] for r in rows}
    flags={r["ValueFlagCode"]:r["ValueFlag"] for r in rows}
    mapping={c:GEOS[v[0]]["id"] for c,v in reporters.items() if v[0] in GEOS}
    # Reviewed source names/codes, not a general inference from numeric prefixes.
    aliases={
        "158":("CHT","Chinese Taipei","TWN"),
        "642":("ROM","Romania","ROU"),
        "254":("","French Guiana","GUF"),
        "312":("","Guadeloupe","GLP"),
        "638":("","Reunion","REU"),
        "474":("","Martinique","MTQ"),
        "535":("","Bonaire, Sint Eustatius and Saba","BES"),
    }
    for code,(source_iso,label,iso) in aliases.items():
        assert reporters[code]==[source_iso,label]
        mapping[code]=GEOS[iso]["id"]
    # Keep historical/aggregate source areas distinct rather than guessing successors.
    extra=json.loads((DATA/"public-geographies.json").read_text("utf-8"))
    extra=[g for g in extra if g["id"] not in {"wto:"+code for code in aliases}]
    for code,(iso,name) in reporters.items():
        if code in mapping: continue
        ident="world" if name=="World" else "wto:"+code
        if ident!="world" and not any(g["id"]==ident for g in extra):
            extra.append(dict(id=ident,label=name+" · WTO",iso3="",regionId="Unassigned",kind="provider_area"))
        mapping[code]=ident
    write(DATA/"public-geographies.json",extra)
    groups=collections.defaultdict(list);titles=collections.defaultdict(dict);seen=set()
    for r in rows:
        assert [r[k] for k in ["PeriodCode","Period","FrequencyCode","Frequency","UnitCode","Unit"]]==["A","Annual","A","Annual","USM","Million US dollar"]
        assert r["ProductClassificationCode"]=="SITC3"
        if r["PartnerCode"]!="000":continue
        assert r["Partner"]=="World" and r["PartnerISO3A"]==""
        c=r["ReporterCode"];pc=r["IndicatorCode"]+":"+r["ProductCode"]
        key=(c,pc,r["Year"]);assert key not in seen;seen.add(key)
        assert 1948<=int(r["Year"])<=2025 and float(r["Value"])>=0
        titles[c][pc]=r["Indicator"]+" · "+r["Product"]
        notes=["Partner: World · SITC Revision 3"]
        if r["ValueFlagCode"]:notes.append(r["ValueFlagCode"]+": "+r["ValueFlag"])
        groups[(c,pc)].append(point(r["Year"],r["Value"],"estimated" if r["ValueFlagCode"]=="E" else "source_statistics",notes))
    areas=[area(c,reporters[c][1],mapping[c],titles[c]) for c in sorted(titles)]
    assert len({a["geographyId"] for a in areas})==len(areas)
    profiles=[profile(metrics[pc],a,sorted(groups[(a["code"],pc)],key=lambda p:p["period"])) for a in areas for pc in sorted(a["seriesTitles"])]
    source(sid,"WTO · Warenhandel nach Produkten","wto_merchandise","wto-values.zip","https://stats.wto.org/assets/UserGuide/merchandise_values_annual_dataset.zip","https://data.wto.org/dataset/bulkdownload","https://www.wto.org/english/res_e/statis_e/about_e.htm","Stand geprüft 15.09.2026","source_statistics",areas,profiles,len(rows))
    AUDIT[-1]['rawSha256']=SOURCES[-1]['expectedSha256']
    SOURCES[-1]['expectedSha256']=AUDIT[-1]['sha256']=wto_fingerprint(rows)
    SOURCES[-1]['recipe']='wto-merchandise-all-csv-fields-multiset-v3'
    AUDIT[-1]['reviewedAliases']={c:dict(sourceIso3=v[0],sourceLabel=v[1],iso3=v[2],geographyId=mapping[c]) for c,v in aliases.items()}
    CONTRACTS[sid]=dict(headers=list(rows[0]),reporters=reporters,products=products,flags=flags,partners={r["PartnerCode"]:[r["PartnerISO3A"],r["Partner"]] for r in rows})

if __name__=="__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    imf();wage();hours();icp();wto()
    ids={s["id"] for s in SOURCES}
    CONFIG["version"]="2026-09-15.1"
    CONFIG["sources"]=[s for s in CONFIG["sources"] if s["id"] not in ids]+SOURCES
    CONFIG["metrics"]=[m for m in CONFIG["metrics"] if m["sourceId"] not in ids]+METRICS
    write(DATA/"public-series-catalog.json",CONFIG)
    write(DATA/"public-gap-contracts.json",CONTRACTS)
    write(HERE/"gap-source-audit-2026-09-15.json",AUDIT)
    print(json.dumps(AUDIT,ensure_ascii=False,indent=2))
